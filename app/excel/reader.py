"""Converte arquivos ``.xlsx`` em snapshots lógicos comparáveis.

O caminho rápido lê diretamente o XML interno do XLSX, evitando a criação de
objetos ``Cell`` do openpyxl. Para estruturas de fórmula não suportadas pelo
leitor direto, há fallback automático para o leitor openpyxl original.
"""

from __future__ import annotations

from io import BytesIO
from dataclasses import dataclass
from collections.abc import Iterator, Mapping
import hashlib
import logging
from pathlib import Path, PurePosixPath
import re
from time import perf_counter
from typing import Callable, TypeAlias
import xml.etree.ElementTree as ET
from xml.sax.saxutils import quoteattr
import zipfile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.datetime import (
    CALENDAR_MAC_1904,
    CALENDAR_WINDOWS_1900,
    from_ISO8601,
    from_excel,
)


logger = logging.getLogger("auditoria_excel.reader")

CellValue: TypeAlias = str | int | float | bool | None
SheetSnapshot: TypeAlias = Mapping[str, CellValue]
Snapshot: TypeAlias = dict[str, SheetSnapshot]


def _member_digest(payload: bytes | None) -> str:
    """SHA-256 que diferencia explicitamente membro ausente de membro vazio."""
    if payload is None:
        return "missing"
    return hashlib.sha256(payload).hexdigest()


@dataclass(frozen=True, slots=True)
class _CachedSheet:
    signature: tuple[str, ...]
    snapshot: SheetSnapshot
    rows: tuple["_CachedRow", ...] = ()


@dataclass(frozen=True, slots=True)
class _CachedRow:
    key: str
    signature: str
    cells: dict[str, CellValue]
    shared_string_indices: tuple[int, ...] = ()
    shared_string_dependencies_supported: bool = True


@dataclass(frozen=True, slots=True)
class SharedStringsSnapshot:
    """Assinaturas por posicao; ``None`` nunca autoriza equivalencia.

    A assinatura cobre o ``<si>`` XML normalizado inteiro, e nao apenas o texto
    exibido. Assim runs de rich text, ``xml:space`` e propriedades foneticas
    nao sao acidentalmente tratados como equivalentes.
    """

    signatures: tuple[str | None, ...]
    values: tuple[str, ...]
    present: bool
    hash_seconds: float = 0.0


class RowSheetSnapshot(Mapping[str, CellValue]):
    """Visão imutável de uma aba composta por mapas imutáveis por row.

    O mapa plano deixa de ser reconstruído a cada versão. Rows cuja identidade
    XML foi provada por SHA-256 compartilham exatamente o mesmo ``dict``.
    """

    __slots__ = ("rows", "_length")

    def __init__(self, rows: tuple[_CachedRow, ...]) -> None:
        self.rows = rows
        self._length = sum(len(row.cells) for row in rows)

    def __getitem__(self, key: str) -> CellValue:
        # Endereços regulares carregam sua row; a busca linear é reservada a
        # acessos pontuais e não participa do comparador incremental.
        for row in self.rows:
            if key in row.cells:
                return row.cells[key]
        raise KeyError(key)

    def __iter__(self) -> Iterator[str]:
        for row in self.rows:
            yield from row.cells

    def __len__(self) -> int:
        return self._length

    def row_maps(self) -> tuple[dict[str, CellValue], ...]:
        return tuple(row.cells for row in self.rows)


_RAW_ROW_START = re.compile(br"<row(?:\s|>)")
_RAW_ROW_NUMBER = re.compile(br"\br=[\"']([^\"']+)[\"']")


def _namespace_wrapper(xml: bytes) -> bytes:
    """Reproduz no wrapper todos os namespaces em escopo na worksheet.

    O digest continua cobrindo exclusivamente os bytes originais da row. O
    wrapper existe apenas para dar ao parser do fragmento o mesmo contexto de
    namespaces que ele teria no documento completo.
    """
    declarations: dict[str, str] = {}
    try:
        for event, value in ET.iterparse(
            BytesIO(xml), events=("start-ns", "start")
        ):
            if event == "start-ns":
                prefix, uri = value
                declarations[prefix] = uri
                continue
            break
    except ET.ParseError as error:
        raise _FastReaderUnsupported(f"XML da worksheet inválido: {error}") from error
    declarations.setdefault("", _MAIN_NS)
    attributes = "".join(
        f" xmlns{':' + prefix if prefix else ''}={quoteattr(uri)}"
        for prefix, uri in declarations.items()
    )
    return f"<root{attributes}>".encode("utf-8")


def _row_element(payload: bytes, namespace_wrapper: bytes) -> ET.Element:
    try:
        return ET.fromstring(namespace_wrapper + payload + b"</root>")[0]
    except ET.ParseError as error:
        raise _FastReaderUnsupported(
            f"fragmento row incompatível com parser incremental: {error}"
        ) from error


def _iter_row_parts(
    xml: bytes, namespace_wrapper: bytes
) -> Iterator[tuple[str, str, bytes, bool]]:
    """Extrai rows e digests fortes sem converter valores de célula.

    O SHA-256 cobre os bytes originais completos da row (atributos, células,
    fórmulas e valores incluídos); o wrapper de namespaces nunca participa do
    digest.
    """
    # SpreadsheetML emitido por Excel/openpyxl usa namespace default e rows
    # não prefixadas. Fora desse formato conservador o chamador abandona o
    # caminho incremental. O digest cobre os bytes XML exatos, não offsets.
    if b"<sheetData" not in xml or re.search(br"<[A-Za-z_][\w.-]*:row(?:\s|>)", xml):
        raise _FastReaderUnsupported("worksheet com rows XML prefixadas")
    occurrence: dict[str, int] = {}
    cursor = 0
    while match := _RAW_ROW_START.search(xml, cursor):
        start = match.start()
        tag_end = xml.find(b">", start)
        if tag_end < 0:
            raise _FastReaderUnsupported("row XML truncada")
        if xml[tag_end - 1:tag_end] == b"/":
            end = tag_end + 1
        else:
            close = xml.find(b"</row>", tag_end + 1)
            if close < 0:
                raise _FastReaderUnsupported("row XML sem fechamento")
            end = close + len(b"</row>")
        payload = xml[start:end]
        number_match = _RAW_ROW_NUMBER.search(payload[: tag_end - start + 1])
        row_number = number_match.group(1).decode("utf-8") if number_match else ""
        ordinal = occurrence.get(row_number, 0)
        occurrence[row_number] = ordinal + 1
        key = f"{row_number}#{ordinal}"
        digest = hashlib.sha256(payload).hexdigest()
        has_shared_formula = False
        if b"<f" in payload:
            element = _row_element(payload, namespace_wrapper)
            for formula in element.iter(_FORMULA_TAG):
                formula_type = formula.get("t")
                if formula_type in {"array", "dataTable"}:
                    raise _FastReaderUnsupported(
                        f"fórmula {formula_type} em row {row_number}"
                    )
                if formula_type == "shared":
                    has_shared_formula = True
        yield key, digest, payload, has_shared_formula
        cursor = end


def _parse_row(
    element: ET.Element,
    shared_strings: list[str],
    date_styles: set[int],
    epoch: object,
) -> tuple[dict[str, CellValue], int, tuple[int, ...], bool]:
    cells: dict[str, CellValue] = {}
    seen = 0
    # Esta função somente é usada quando a planilha não contém fórmula
    # compartilhada, logo o estado de tradução nunca cruza fronteiras de row.
    formulas: dict[str, tuple[str, str]] = {}
    shared_indices: set[int] = set()
    dependencies_supported = True
    for cell in element.iter(_tag("c")):
        seen += 1
        coordinate = cell.get("r")
        if coordinate:
            if cell.get("t") == "s" and cell.find("m:f", _NS) is None:
                value_element = cell.find("m:v", _NS)
                try:
                    if value_element is None or value_element.text is None:
                        dependencies_supported = False
                    else:
                        index = int(value_element.text)
                        if index < 0:
                            dependencies_supported = False
                        else:
                            shared_indices.add(index)
                except ValueError:
                    dependencies_supported = False
            value = _cell_value(
                cell, coordinate, shared_strings, date_styles, epoch, formulas
            )
            if value is not None:
                cells[coordinate] = value  # type: ignore[assignment]
    return cells, seen, tuple(sorted(shared_indices)), dependencies_supported


@dataclass(frozen=True, slots=True)
class ReaderMetrics:
    worksheets: int
    worksheets_reused: int
    cells_parsed: int
    fallback_used: bool
    rows_total: int = 0
    rows_reused: int = 0
    rows_parsed: int = 0
    cells_reused: int = 0
    hash_seconds: float = 0.0
    dependency_seconds: float = 0.0
    structural_diff_seconds: float = 0.0
    parsing_seconds: float = 0.0
    snapshot_seconds: float = 0.0
    fallback_reason: str | None = None
    sharedstrings_total_previous: int = 0
    sharedstrings_total_current: int = 0
    sharedstrings_indices_equal: int = 0
    sharedstrings_indices_changed: int = 0
    sharedstrings_indices_new: int = 0
    sharedstrings_indices_removed: int = 0
    sharedstrings_shadow_rows_candidate: int = 0
    sharedstrings_shadow_rows_safe: int = 0
    sharedstrings_shadow_rows_invalidated: int = 0
    sharedstrings_shadow_indices_checked: int = 0
    sharedstrings_shadow_indices_changed: int = 0
    sharedstrings_shadow_indices_new: int = 0
    rows_dependent_sharedstrings: int = 0
    rows_reused_sharedstrings: int = 0
    rows_invalidated_changed_index: int = 0
    rows_invalidated_global_dependency: int = 0
    rows_invalidated_unsupported_structure: int = 0
    sharedstrings_diff_seconds: float = 0.0
    sharedstrings_hash_seconds: float = 0.0

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_NS = {"m": _MAIN_NS, "r": _REL_NS, "pr": _PKG_REL_NS}


class _FastReaderUnsupported(RuntimeError):
    """Sinaliza recurso raro em que o fallback openpyxl preserva melhor a semântica."""


def _tag(local: str) -> str:
    return f"{{{_MAIN_NS}}}{local}"


_FORMULA_TAG = _tag("f")
_VALUE_TAG = _tag("v")
_INLINE_STRING_TAG = _tag("is")


def _read_openpyxl(path: Path) -> Snapshot:
    """Fallback equivalente ao leitor original."""
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        snapshot: Snapshot = {}
        for worksheet in workbook.worksheets:
            cells: SheetSnapshot = {}
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells[cell.coordinate] = cell.value
            snapshot[worksheet.title] = cells
        return snapshot
    finally:
        workbook.close()


def _cast_number(value: str) -> int | float:
    text = value.strip()
    if not text:
        raise ValueError("valor numérico vazio")
    if all(char not in text for char in ".eE"):
        try:
            return int(text)
        except ValueError:
            pass
    return float(text)


def _all_text(element: ET.Element | None) -> str:
    if element is None:
        return ""
    return "".join(node.text or "" for node in element.iter() if node.tag == _tag("t"))


def _shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    result: list[str] = []
    with archive.open("xl/sharedStrings.xml") as stream:
        for event, element in ET.iterparse(stream, events=("end",)):
            if element.tag == _tag("si"):
                result.append(_all_text(element))
                element.clear()
    return result


_SUPPORTED_SHARED_STRING_ELEMENTS = {
    _tag(name)
    for name in (
        "si", "t", "r", "rPr", "rFont", "charset", "family", "b", "i",
        "strike", "outline", "shadow", "condense", "extend", "color", "sz",
        "u", "vertAlign", "scheme", "rPh", "phoneticPr",
    )
}


def _shared_strings_snapshot(payload: bytes | None) -> SharedStringsSnapshot:
    """Constroi prova conservadora por indice a partir do XML completo."""
    started = perf_counter()
    if payload is None:
        return SharedStringsSnapshot((), (), False, perf_counter() - started)
    try:
        root = ET.fromstring(payload)
    except ET.ParseError as error:
        raise _FastReaderUnsupported(f"sharedStrings XML invalido: {error}") from error
    if root.tag != _tag("sst"):
        raise _FastReaderUnsupported("raiz de sharedStrings nao suportada")
    signatures: list[str | None] = []
    values: list[str] = []
    for entry in root:
        if entry.tag != _tag("si"):
            # Metadado no nivel sst nao altera os indices, mas torna a prova
            # desconhecida; nenhuma entrada e autorizada nesse documento.
            return SharedStringsSnapshot(
                tuple(None for _ in root.findall("m:si", _NS)),
                tuple(_all_text(si) for si in root.findall("m:si", _NS)),
                True,
                perf_counter() - started,
            )
        values.append(_all_text(entry))
        supported = all(
            node.tag in _SUPPORTED_SHARED_STRING_ELEMENTS for node in entry.iter()
        )
        if not supported:
            signatures.append(None)
            continue
        # ElementTree ja normaliza entidades/Unicode ao interpretar. Serializar
        # a arvore inteira preserva tags, atributos, texto e tails relevantes.
        # Diferencas cosmeticas que ele nao normaliza geram apenas falso
        # negativo (parse normal), nunca autorizacao indevida. C14N por entrada
        # mostrou-se mais caro que o parsing que este shadow mode mede.
        semantic_xml = ET.tostring(entry, encoding="utf-8")
        signatures.append(hashlib.sha256(semantic_xml).hexdigest())
    return SharedStringsSnapshot(
        tuple(signatures), tuple(values), True, perf_counter() - started
    )


def _workbook_epoch(archive: zipfile.ZipFile) -> object:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    workbook_pr = root.find("m:workbookPr", _NS)
    if workbook_pr is not None and workbook_pr.get("date1904") in {"1", "true", "True"}:
        return CALENDAR_MAC_1904
    return CALENDAR_WINDOWS_1900


def _date_styles(archive: zipfile.ZipFile) -> set[int]:
    if "xl/styles.xml" not in archive.namelist():
        return set()
    root = ET.fromstring(archive.read("xl/styles.xml"))
    custom_formats: dict[int, str] = {}
    num_fmts = root.find("m:numFmts", _NS)
    if num_fmts is not None:
        for num_fmt in num_fmts.findall("m:numFmt", _NS):
            try:
                num_fmt_id = int(num_fmt.get("numFmtId", ""))
            except ValueError:
                continue
            format_code = num_fmt.get("formatCode")
            if format_code is not None:
                custom_formats[num_fmt_id] = format_code

    result: set[int] = set()
    cell_xfs = root.find("m:cellXfs", _NS)
    if cell_xfs is None:
        return result
    for index, xf in enumerate(cell_xfs.findall("m:xf", _NS)):
        try:
            num_fmt_id = int(xf.get("numFmtId", "0"))
        except ValueError:
            num_fmt_id = 0
        format_code = custom_formats.get(num_fmt_id) or BUILTIN_FORMATS.get(num_fmt_id)
        if format_code and is_date_format(format_code):
            result.add(index)
    return result


def _sheet_targets(archive: zipfile.ZipFile) -> list[tuple[str, str]]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets: dict[str, str] = {}
    for rel in rels.findall("pr:Relationship", _NS):
        rel_id = rel.get("Id")
        target = rel.get("Target")
        if rel_id and target:
            rel_targets[rel_id] = target

    result: list[tuple[str, str]] = []
    sheets = workbook.find("m:sheets", _NS)
    if sheets is None:
        return result
    for sheet in sheets.findall("m:sheet", _NS):
        title = sheet.get("name")
        rel_id = sheet.get(f"{{{_REL_NS}}}id")
        if not title or not rel_id or rel_id not in rel_targets:
            continue
        target = rel_targets[rel_id].lstrip("/")
        if not target.startswith("xl/"):
            target = str(PurePosixPath("xl") / target)
        result.append((title, target))
    return result


def _formula_element_value(
    formula: ET.Element | None,
    coordinate: str,
    shared_formulas: dict[str, tuple[str, str]],
) -> str | None:
    if formula is None:
        return None
    formula_type = formula.get("t")
    if formula_type in {"array", "dataTable"}:
        raise _FastReaderUnsupported(f"fórmula {formula_type} em {coordinate}")

    text = formula.text or ""
    if formula_type == "shared":
        shared_index = formula.get("si")
        if shared_index is None:
            raise _FastReaderUnsupported(f"fórmula compartilhada sem si em {coordinate}")
        if text:
            value = "=" + text
            shared_formulas[shared_index] = (coordinate, value)
            return value
        master = shared_formulas.get(shared_index)
        if master is None:
            raise _FastReaderUnsupported(
                f"fórmula compartilhada sem mestre conhecido em {coordinate}"
            )
        origin, master_formula = master
        return Translator(master_formula, origin=origin).translate_formula(coordinate)

    return "=" + text


def _formula_value(
    cell: ET.Element,
    coordinate: str,
    shared_formulas: dict[str, tuple[str, str]],
) -> str | None:
    """Compatibility helper used by the profiling baseline."""
    return _formula_element_value(cell.find("m:f", _NS), coordinate, shared_formulas)


def _convert_cell_value(
    cell: ET.Element,
    coordinate: str,
    formula_element: ET.Element | None,
    value_element: ET.Element | None,
    inline_element: ET.Element | None,
    shared_strings: list[str],
    date_styles: set[int],
    epoch: object,
    shared_formulas: dict[str, tuple[str, str]],
) -> object:
    formula = _formula_element_value(formula_element, coordinate, shared_formulas)
    if formula is not None:
        return formula

    cell_type = cell.get("t", "n")
    raw = value_element.text if value_element is not None else None

    if cell_type == "inlineStr":
        text = _all_text(inline_element)
        return text if text != "" else None
    if raw is None:
        return None
    if cell_type == "s":
        try:
            index = int(raw)
            if index < 0:
                raise IndexError(index)
            return shared_strings[index]
        except (ValueError, IndexError) as error:
            raise _FastReaderUnsupported(
                f"indice sharedStrings invalido em {coordinate}: {raw!r}"
            ) from error
    if cell_type == "b":
        return raw == "1"
    if cell_type in {"str", "e"}:
        return raw
    if cell_type == "d":
        return from_ISO8601(raw)

    value = _cast_number(raw)
    style_text = cell.get("s")
    if style_text is not None:
        try:
            style_id = int(style_text)
        except ValueError:
            style_id = -1
        if style_id in date_styles:
            return from_excel(value, epoch)
    return value


def _cell_value(
    cell: ET.Element,
    coordinate: str,
    shared_strings: list[str],
    date_styles: set[int],
    epoch: object,
    shared_formulas: dict[str, tuple[str, str]],
) -> object:
    """Converte uma célula após uma única visita a seus filhos XML diretos."""
    formula_element = value_element = inline_element = None
    for child in cell:
        # Element.find() selecionava a primeira ocorrência; mantenha a mesma
        # semântica até para XML incomum com filhos duplicados.
        if child.tag == _FORMULA_TAG and formula_element is None:
            formula_element = child
        elif child.tag == _VALUE_TAG and value_element is None:
            value_element = child
        elif child.tag == _INLINE_STRING_TAG and inline_element is None:
            inline_element = child
    return _convert_cell_value(
        cell,
        coordinate,
        formula_element,
        value_element,
        inline_element,
        shared_strings,
        date_styles,
        epoch,
        shared_formulas,
    )


def _cell_value_repeated_find(
    cell: ET.Element,
    coordinate: str,
    shared_strings: list[str],
    date_styles: set[int],
    epoch: object,
    shared_formulas: dict[str, tuple[str, str]],
) -> object:
    """Implementação anterior, mantida apenas como baseline de equivalência."""
    formula = _formula_value(cell, coordinate, shared_formulas)
    if formula is not None:
        return formula

    cell_type = cell.get("t", "n")
    value_element = cell.find("m:v", _NS)
    raw = value_element.text if value_element is not None else None
    if cell_type == "inlineStr":
        text = _all_text(cell.find("m:is", _NS))
        return text if text != "" else None
    if raw is None:
        return None
    if cell_type == "s":
        return shared_strings[int(raw)]
    if cell_type == "b":
        return raw == "1"
    if cell_type in {"str", "e"}:
        return raw
    if cell_type == "d":
        return from_ISO8601(raw)

    value = _cast_number(raw)
    style_text = cell.get("s")
    if style_text is not None:
        try:
            style_id = int(style_text)
        except ValueError:
            style_id = -1
        if style_id in date_styles:
            return from_excel(value, epoch)
    return value


def _read_fast_with(
    path: Path,
    cell_reader: Callable[
        [ET.Element, str, list[str], set[int], object, dict[str, tuple[str, str]]],
        object,
    ],
) -> Snapshot:
    with zipfile.ZipFile(path) as archive:
        shared = _shared_strings(archive)
        date_styles = _date_styles(archive)
        epoch = _workbook_epoch(archive)
        sheets = _sheet_targets(archive)
        snapshot: Snapshot = {}

        for title, target in sheets:
            cells: SheetSnapshot = {}
            shared_formulas: dict[str, tuple[str, str]] = {}
            with archive.open(target) as stream:
                for event, element in ET.iterparse(stream, events=("end",)):
                    if element.tag != _tag("c"):
                        continue
                    coordinate = element.get("r")
                    if coordinate:
                        value = cell_reader(
                            element,
                            coordinate,
                            shared,
                            date_styles,
                            epoch,
                            shared_formulas,
                        )
                        if value is not None:
                            cells[coordinate] = value  # type: ignore[assignment]
                    element.clear()
            snapshot[title] = cells
        return snapshot


def _read_fast(path: Path) -> Snapshot:
    return _read_fast_with(path, _cell_value)


def _read_fast_repeated_find(path: Path) -> Snapshot:
    """Executa o leitor anterior para testes e benchmarks, nunca em produção."""
    return _read_fast_with(path, _cell_value_repeated_find)


class ConsecutiveWorkbookReader:
    """Lê versões consecutivas reutilizando somente abas provadas idênticas.

    A identidade compartilhada de ``SheetSnapshot`` é concedida apenas quando
    SHA-256 da worksheet e de todas as dependências usadas na interpretação
    coincide. O cache contém exclusivamente a última versão, limitando memória.
    """

    def __init__(self) -> None:
        self._previous: dict[str, _CachedSheet] = {}
        self._previous_shared = SharedStringsSnapshot((), (), False)
        self.last_metrics = ReaderMetrics(0, 0, 0, False)

    def read(self, path: str | Path) -> Snapshot:
        workbook_path = Path(path)
        if workbook_path.suffix.lower() != ".xlsx":
            raise ValueError("O leitor aceita somente arquivos .xlsx")
        try:
            snapshot, cache, metrics, shared_snapshot = self._read_fast(workbook_path)
        except _FastReaderUnsupported as error:
            # Um fallback é deliberadamente uma barreira de cache: snapshots
            # openpyxl não receberam a prova criptográfica desta classe.
            self._previous = {}
            self._previous_shared = SharedStringsSnapshot((), (), False)
            logger.info(
                "Leitor incremental usou fallback openpyxl arquivo=%s motivo=%s",
                workbook_path.name,
                error,
            )
            snapshot = _read_openpyxl(workbook_path)
            self.last_metrics = ReaderMetrics(
                len(snapshot), 0, 0, True, fallback_reason=str(error)
            )
            return snapshot
        self._previous = cache
        self._previous_shared = shared_snapshot
        self.last_metrics = metrics
        logger.debug(
            "Leitura XLSX incremental arquivo=%s abas=%d reutilizadas=%d celulas_parseadas=%d",
            workbook_path.name,
            metrics.worksheets,
            metrics.worksheets_reused,
            metrics.cells_parsed,
        )
        return snapshot

    def _read_fast(
        self, path: Path
    ) -> tuple[
        Snapshot, dict[str, _CachedSheet], ReaderMetrics, SharedStringsSnapshot
    ]:
        with zipfile.ZipFile(path) as archive:
            dependency_started = perf_counter()
            names = set(archive.namelist())

            def member(name: str) -> bytes | None:
                return archive.read(name) if name in names else None

            shared_xml = member("xl/sharedStrings.xml")
            styles_xml = member("xl/styles.xml")
            relationships_xml = member("xl/_rels/workbook.xml.rels")
            # Estas rotinas conservam exatamente a conversão do leitor oficial.
            shared_snapshot = _shared_strings_snapshot(shared_xml)
            shared = list(shared_snapshot.values)
            date_styles = _date_styles(archive)
            epoch = _workbook_epoch(archive)
            sheets = _sheet_targets(archive)
            dependency_signature = (
                _member_digest(shared_xml),
                _member_digest(styles_xml),
                "1904" if epoch == CALENDAR_MAC_1904 else "1900",
                _member_digest(relationships_xml),
            )
            dependency_seconds = perf_counter() - dependency_started
            diff_started = perf_counter()
            previous_shared = self._previous_shared
            common = min(len(previous_shared.signatures), len(shared_snapshot.signatures))
            shared_equal = sum(
                1 for index in range(common)
                if previous_shared.signatures[index] is not None
                and previous_shared.signatures[index] == shared_snapshot.signatures[index]
            )
            shared_changed = common - shared_equal
            shared_new = max(0, len(shared_snapshot.signatures) - common)
            shared_removed = max(0, len(previous_shared.signatures) - common)
            shared_diff_seconds = perf_counter() - diff_started
            snapshot: Snapshot = {}
            new_cache: dict[str, _CachedSheet] = {}
            reused = 0
            parsed = 0
            rows_total = rows_reused = rows_parsed = cells_reused = 0
            hash_seconds = structural_seconds = parsing_seconds = snapshot_seconds = 0.0
            shadow_candidate = shadow_safe = shadow_invalidated = 0
            shadow_checked = shadow_changed = shadow_new = 0
            rows_dependent = rows_reused_shared = invalid_changed = 0
            invalid_global = invalid_unsupported = 0
            for title, target in sheets:
                worksheet_xml = archive.read(target)
                hash_started = perf_counter()
                signature = (
                    _member_digest(worksheet_xml),
                    *dependency_signature,
                    target,
                )
                hash_seconds += perf_counter() - hash_started
                cached = self._previous.get(target)
                if cached is not None and cached.signature == signature:
                    cells = cached.snapshot
                    reused += 1
                    rows_total += len(cached.rows)
                    rows_reused += len(cached.rows)
                    cells_reused += len(cells)
                else:
                    diff_started = perf_counter()
                    parsing_before_diff = parsing_seconds
                    namespace_wrapper = _namespace_wrapper(worksheet_xml)
                    dependencies_equal = (
                        cached is not None and cached.signature[1:] == signature[1:]
                    )
                    dependencies_except_shared_equal = (
                        cached is not None and cached.signature[2:] == signature[2:]
                    )
                    previous_rows = {
                        (row.key, row.signature): row
                        for row in (cached.rows if dependencies_equal else ())
                    }
                    shadow_previous_rows = {
                        (row.key, row.signature): row
                        for row in (cached.rows if dependencies_except_shared_equal else ())
                    }
                    built_rows: list[_CachedRow] = []
                    sheet_rows = sheet_reused = sheet_parsed = sheet_cells_reused = 0
                    has_shared_formula = False
                    force_full_parse = False
                    for key, row_digest, row_xml, row_has_shared in _iter_row_parts(
                        worksheet_xml, namespace_wrapper
                    ):
                        sheet_rows += 1
                        has_shared_formula = has_shared_formula or row_has_shared
                        if has_shared_formula:
                            continue
                        old_row = previous_rows.get((key, row_digest))
                        if old_row is not None:
                            built_rows.append(old_row)
                            sheet_reused += 1
                            sheet_cells_reused += len(old_row.cells)
                            continue
                        parse_started = perf_counter()
                        row_cells, seen, shared_indices, deps_supported = _parse_row(
                            _row_element(row_xml, namespace_wrapper),
                            shared,
                            date_styles,
                            epoch,
                        )
                        parsing_seconds += perf_counter() - parse_started
                        parsed += seen
                        sheet_parsed += 1
                        built_rows.append(_CachedRow(
                            key, row_digest, row_cells, shared_indices, deps_supported
                        ))
                        if shared_indices:
                            rows_dependent += 1
                        shadow_old = shadow_previous_rows.get((key, row_digest))
                        shared_globally_changed = (
                            cached is not None and cached.signature[1] != signature[1]
                        )
                        if shadow_old is not None and shared_globally_changed:
                            shadow_candidate += 1
                            if shadow_old.shared_string_indices:
                                invalid_global += 1
                            safe = shadow_old.shared_string_dependencies_supported
                            if not safe:
                                invalid_unsupported += 1
                            for index in shadow_old.shared_string_indices:
                                shadow_checked += 1
                                if index >= len(shared_snapshot.signatures):
                                    shadow_new += 1
                                    safe = False
                                elif index >= len(previous_shared.signatures):
                                    shadow_new += 1
                                    safe = False
                                elif previous_shared.signatures[index] is None or (
                                    previous_shared.signatures[index]
                                    != shared_snapshot.signatures[index]
                                ):
                                    shadow_changed += 1
                                    invalid_changed += 1
                                    safe = False
                            # Validacao sombra obrigatoria contra o resultado
                            # recem-parseado pelo caminho oficial atual.
                            safe = safe and shadow_old.cells == row_cells
                            if safe:
                                shadow_safe += 1
                                if shadow_old.shared_string_indices:
                                    rows_reused_shared += 1
                            else:
                                shadow_invalidated += 1
                        if (
                            dependencies_equal
                            and sheet_rows == 256
                            and sheet_reused < 26
                        ):
                            force_full_parse = True
                            break
                    structural_seconds += max(
                        0.0,
                        perf_counter() - diff_started
                        - (parsing_seconds - parsing_before_diff),
                    )
                    if force_full_parse:
                        sheet_rows = sum(1 for _ in _RAW_ROW_START.finditer(worksheet_xml))
                    rows_total += sheet_rows
                    # Shared formulas carry a master/follower state across rows.
                    # Until that dependency is indexed explicitly, retain the
                    # proven full-sheet parser rather than risk stale formulas.
                    if has_shared_formula or force_full_parse:
                        parse_started = perf_counter()
                        plain: dict[str, CellValue] = {}
                        shared_formulas: dict[str, tuple[str, str]] = {}
                        for _, element in ET.iterparse(
                            BytesIO(worksheet_xml), events=("end",)
                        ):
                            if element.tag != _tag("c"):
                                continue
                            parsed += 1
                            coordinate = element.get("r")
                            if coordinate:
                                value = _cell_value(
                                    element, coordinate, shared, date_styles, epoch,
                                    shared_formulas,
                                )
                                if value is not None:
                                    plain[coordinate] = value  # type: ignore[assignment]
                            element.clear()
                        parsing_seconds += perf_counter() - parse_started
                        rows_parsed += sheet_rows
                        cells = plain
                        cached_rows: tuple[_CachedRow, ...] = ()
                    else:
                        rows_reused += sheet_reused
                        rows_parsed += sheet_parsed
                        cells_reused += sheet_cells_reused
                        snapshot_started = perf_counter()
                        cached_rows = tuple(built_rows)
                        cells = RowSheetSnapshot(cached_rows)
                        snapshot_seconds += perf_counter() - snapshot_started
                snapshot[title] = cells
                new_cache[target] = _CachedSheet(
                    signature, cells,
                    cached.rows if cached is not None and cached.signature == signature
                    else cached_rows,
                )
            metrics = ReaderMetrics(
                len(sheets), reused, parsed, False, rows_total, rows_reused,
                rows_parsed, cells_reused, hash_seconds, dependency_seconds,
                structural_seconds, parsing_seconds, snapshot_seconds, None,
                len(previous_shared.signatures), len(shared_snapshot.signatures),
                shared_equal, shared_changed, shared_new, shared_removed,
                shadow_candidate, shadow_safe, shadow_invalidated,
                shadow_checked, shadow_changed, shadow_new,
                rows_dependent, rows_reused_shared, invalid_changed,
                invalid_global, invalid_unsupported, shared_diff_seconds,
                shared_snapshot.hash_seconds,
            )
            return snapshot, new_cache, metrics, shared_snapshot


def read_workbook(path: str | Path) -> Snapshot:
    """Lê um ``.xlsx`` e preserva a semântica usada pela auditoria.

    O leitor XML é usado por padrão. Recursos raros cuja equivalência não pode
    ser garantida caem automaticamente no openpyxl original.
    """
    workbook_path = Path(path)
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("O leitor aceita somente arquivos .xlsx")

    try:
        snapshot = _read_fast(workbook_path)
        logger.debug("Leitura XLSX concluída modo=xml arquivo=%s", workbook_path.name)
        return snapshot
    except _FastReaderUnsupported as error:
        logger.info(
            "Leitor XML usou fallback openpyxl arquivo=%s motivo=%s",
            workbook_path.name,
            error,
        )
        return _read_openpyxl(workbook_path)
