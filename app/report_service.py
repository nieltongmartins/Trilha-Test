"""Exportação streaming, íntegra e atômica da trilha persistida no SQLite."""

from __future__ import annotations

from collections.abc import Callable, Iterator, Sequence
from datetime import datetime, timezone
from pathlib import Path
import logging
import os
import sqlite3
import time
import uuid
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.report_artifacts import ReportArtifactManager


logger = logging.getLogger("auditoria_excel.report")

CHANGE_ROWS_PER_SHEET = 900_000
FETCH_BATCH_SIZE = 10_000
EXCEL_MAX_ROWS = 1_048_576

CHANGE_HEADERS = (
    "ID", "Versão anterior", "Versão atual", "Data/hora", "Autor",
    "Comentário", "Aba", "Célula", "Tipo", "Valor anterior", "Valor novo",
)


class ReportValidationError(RuntimeError):
    """Indica que um XLSX temporário não representa integralmente o banco."""


class ReportService:
    """Consulta o SQLite em lotes e só publica um XLSX completamente validado."""

    def __init__(
        self,
        connection: sqlite3.Connection,
        output_directory: str | Path,
        *,
        progress_callback: Callable[[str], None] | None = None,
        change_rows_per_sheet: int = CHANGE_ROWS_PER_SHEET,
        fetch_batch_size: int = FETCH_BATCH_SIZE,
    ) -> None:
        if not 1 <= change_rows_per_sheet < EXCEL_MAX_ROWS:
            raise ValueError("Limite de alterações por aba inválido")
        if fetch_batch_size < 1:
            raise ValueError("Tamanho de lote inválido")
        self.connection = connection
        self.output_directory = Path(output_directory)
        self.progress_callback = progress_callback
        self.change_rows_per_sheet = change_rows_per_sheet
        self.fetch_batch_size = fetch_batch_size
        self._last_progress = 0.0

    def generate(self, spreadsheet_id: int) -> Path:
        started = time.monotonic()
        export_id = str(uuid.uuid4())
        exported_at = datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")
        self._progress("Consultando banco...", force=True)
        spreadsheet = self.connection.execute(
            "SELECT * FROM planilha WHERE id = ?", (spreadsheet_id,)
        ).fetchone()
        if spreadsheet is None:
            raise ValueError("Planilha não encontrada no banco de auditoria")
        stats = self._database_stats(spreadsheet_id)
        self._progress(
            f"Total encontrado: {stats['changes']:,} alterações", force=True
        )

        self.output_directory.mkdir(parents=True, exist_ok=True)
        artifacts = ReportArtifactManager(self.connection, self.output_directory)
        output = artifacts.canonical_path(artifacts.identity(spreadsheet_id))
        temporary = output.with_name(f".{output.stem}.{export_id}.part.xlsx")
        temporary.unlink(missing_ok=True)
        try:
            exported = self._write_workbook(
                temporary, spreadsheet_id, spreadsheet, stats, export_id, exported_at
            )
            self._progress("Validando integridade...", force=True)
            self._validate(temporary, stats, exported)
            os.replace(temporary, output)
        except Exception:
            temporary.unlink(missing_ok=True)
            logger.exception("Exportação FAILED; relatório anterior foi preservado")
            raise

        self._progress("Relatório concluído", force=True)
        logger.info(
            "Relatório gerado e publicado planilha_id=%d alteracoes=%d abas=%d tempo=%.2fs arquivo=%s",
            spreadsheet_id, exported["changes"], len(exported["sheets"]),
            time.monotonic() - started, output,
        )
        return output

    def _database_stats(self, spreadsheet_id: int) -> dict[str, object]:
        changes = self.connection.execute(
            """SELECT COUNT(*) AS total,
                      COALESCE(SUM(tipo='ADD'), 0) AS adds,
                      COALESCE(SUM(tipo='MOD'), 0) AS mods,
                      COALESCE(SUM(tipo='DEL'), 0) AS dels,
                      MIN(id) AS first_id, MAX(id) AS last_id
                 FROM alteracao WHERE planilha_id=?""", (spreadsheet_id,)
        ).fetchone()
        scalar = lambda table: self.connection.execute(  # noqa: E731
            f"SELECT COUNT(*) FROM {table} WHERE planilha_id=?", (spreadsheet_id,)
        ).fetchone()[0]
        checkpoint = self.connection.execute(
            "SELECT versao_numero FROM checkpoint WHERE planilha_id=?", (spreadsheet_id,)
        ).fetchone()
        return {
            "changes": changes["total"], "ADD": changes["adds"],
            "MOD": changes["mods"], "DEL": changes["dels"],
            "first_id": changes["first_id"], "last_id": changes["last_id"],
            "versions": scalar("versao_processada"),
            "executions": scalar("execucao_auditoria"),
            "errors": scalar("erro_processamento"),
            "checkpoint": checkpoint[0] if checkpoint else None,
        }

    def _write_workbook(self, path: Path, spreadsheet_id: int, spreadsheet: sqlite3.Row,
                        stats: dict[str, object], export_id: str,
                        exported_at: str) -> dict[str, object]:
        workbook = Workbook(write_only=True)
        identity = "/".join((spreadsheet["site_id"], spreadsheet["drive_id"],
                             spreadsheet["drive_item_id"]))
        summary = workbook.create_sheet("Resumo")
        self._setup_sheet(summary, (24, 60))
        self._header(summary, ("Campo", "Valor"))
        for row in (
            ("Planilha", spreadsheet["nome_atual"]), ("Identidade técnica", identity),
            ("DriveItem ID", spreadsheet["drive_item_id"]),
            ("Versões processadas", stats["versions"]),
            ("Total de alterações", stats["changes"]), ("ADD", stats["ADD"]),
            ("MOD", stats["MOD"]), ("DEL", stats["DEL"]),
            ("Execuções", stats["executions"]), ("Erros", stats["errors"]),
            ("Checkpoint", stats["checkpoint"]),
        ):
            summary.append(row)
        summary.auto_filter.ref = "A1:B12"

        self._progress("Exportando execuções...", force=True)
        executions = self._export_query(workbook, "Execuções", (
            "ID", "Código", "Início", "Fim", "Checkpoint inicial", "Versão final",
            "Versões", "Alterações", "Status", "Mensagem",
        ), """SELECT id, codigo_execucao, inicio, fim, checkpoint_inicial, versao_final,
                     versoes_processadas, alteracoes_encontradas, status, mensagem
                FROM execucao_auditoria WHERE planilha_id=? ORDER BY id""",
            (spreadsheet_id,))
        self._progress("Exportando versões...", force=True)
        versions = self._export_query(workbook, "Versões processadas", (
            "ID", "Versão anterior", "Versão atual", "Data/hora", "Autor", "E-mail",
            "Login", "Comentário", "Status", "Alterações", "Processamento",
        ), """SELECT id, versao_anterior_numero, versao_atual_numero,
                     data_hora_versao, autor, autor_email, autor_login, comentario,
                     status, quantidade_alteracoes, data_processamento
                FROM versao_processada WHERE planilha_id=? ORDER BY id""",
            (spreadsheet_id,))

        change_result = self._export_changes(workbook, spreadsheet_id, int(stats["changes"]))
        self._progress("Exportando erros...", force=True)
        errors = self._export_query(workbook, "Erros", (
            "ID", "Execução", "Versão anterior", "Versão atual", "Tipo", "Mensagem",
            "Data/hora",
        ), """SELECT id, execucao_id, versao_anterior, versao_atual, tipo_erro,
                     mensagem, data_hora FROM erro_processamento
                WHERE planilha_id=? ORDER BY id""", (spreadsheet_id,))

        result: dict[str, object] = {
            "versions": versions, "executions": executions, "errors": errors,
            **change_result,
        }
        integrity = workbook.create_sheet("Integridade")
        self._setup_sheet(integrity, (34, 80))
        self._header(integrity, ("Campo", "Valor"))
        integrity_rows = [
            ("Identificação técnica", identity), ("Nome da planilha", spreadsheet["nome_atual"]),
            ("Data/hora da exportação", exported_at), ("ID único da exportação", export_id),
            ("Status da exportação", "CONCLUÍDA"),
            ("Versões no banco", stats["versions"]), ("Versões exportadas", versions),
            ("Alterações no banco", stats["changes"]), ("Alterações exportadas", result["changes"]),
            ("ADD no banco", stats["ADD"]), ("MOD no banco", stats["MOD"]),
            ("DEL no banco", stats["DEL"]), ("ADD exportados", result["ADD"]),
            ("MOD exportados", result["MOD"]), ("DEL exportados", result["DEL"]),
            ("Número de abas de alterações", len(result["sheets"])),
            ("Primeiro ID global", stats["first_id"]), ("Último ID global", stats["last_id"]),
            ("Checkpoint observado", stats["checkpoint"]),
            ("Resultado da validação", "VALIDADO"),
        ]
        for sheet in result["sheets"]:
            integrity_rows.extend((
                (f"Registros em {sheet['name']}", sheet["count"]),
                (f"Primeiro ID em {sheet['name']}", sheet["first_id"]),
                (f"Último ID em {sheet['name']}", sheet["last_id"]),
            ))
        for row in integrity_rows:
            integrity.append(row)
        integrity.auto_filter.ref = f"A1:B{len(integrity_rows) + 1}"
        self._progress("Finalizando arquivo...", force=True)
        workbook.save(path)
        return result

    def _export_changes(self, workbook: Workbook, spreadsheet_id: int,
                        total: int) -> dict[str, object]:
        cursor = self.connection.execute(
            """SELECT a.id, v.versao_anterior_numero, v.versao_atual_numero,
                      v.data_hora_versao, v.autor, v.comentario, a.aba, a.endereco,
                      a.tipo, a.valor_anterior, a.valor_novo
                 FROM alteracao a JOIN versao_processada v
                   ON v.id=a.versao_processada_id
                WHERE a.planilha_id=? ORDER BY a.id""", (spreadsheet_id,)
        )
        sheets: list[dict[str, object]] = []
        sheet = None
        current: dict[str, object] | None = None
        exported = 0
        counts = {"ADD": 0, "MOD": 0, "DEL": 0}
        previous_id: int | None = None
        for batch in self._batches(cursor):
            for row in batch:
                if sheet is None or current is None or current["count"] == self.change_rows_per_sheet:
                    name = f"Alterações_{len(sheets) + 1:03d}"
                    self._progress(f"Gerando {name}", force=True)
                    sheet = workbook.create_sheet(name)
                    self._setup_sheet(sheet, (12, 16, 16, 23, 24, 35, 24, 14, 10, 35, 35))
                    self._header(sheet, CHANGE_HEADERS)
                    current = {"name": name, "count": 0, "first_id": None, "last_id": None}
                    sheets.append(current)
                change_id = row[0]
                if previous_id is not None and change_id <= previous_id:
                    raise ReportValidationError("IDs duplicados ou ordem global violada")
                previous_id = change_id
                sheet.append(tuple(row))
                current["first_id"] = current["first_id"] or change_id
                current["last_id"] = change_id
                current["count"] = int(current["count"]) + 1
                exported += 1
                counts[row[8]] += 1
            self._progress(f"Exportando {exported:,} / {total:,}")
        if not sheets:
            name = "Alterações_001"
            sheet = workbook.create_sheet(name)
            self._setup_sheet(sheet, (12,) * len(CHANGE_HEADERS))
            self._header(sheet, CHANGE_HEADERS)
            sheets.append({"name": name, "count": 0, "first_id": None, "last_id": None})
        for item in sheets:
            # Cabeçalho + dados; o filtro não exige materializar células.
            workbook[item["name"]].auto_filter.ref = (
                f"A1:K{int(item['count']) + 1}"
            )
        return {"changes": exported, **counts, "sheets": sheets}

    def _export_query(self, workbook: Workbook, name: str, headers: Sequence[str],
                      sql: str, parameters: tuple[object, ...]) -> int:
        sheet = workbook.create_sheet(name)
        self._setup_sheet(sheet, (18,) * len(headers))
        self._header(sheet, headers)
        count = 0
        for batch in self._batches(self.connection.execute(sql, parameters)):
            for row in batch:
                sheet.append(tuple(row))
                count += 1
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{count + 1}"
        return count

    def _batches(self, cursor: sqlite3.Cursor) -> Iterator[Sequence[sqlite3.Row]]:
        while batch := cursor.fetchmany(self.fetch_batch_size):
            yield batch

    @staticmethod
    def _setup_sheet(sheet, widths: Sequence[int]) -> None:
        sheet.freeze_panes = "A2"
        for index, width in enumerate(widths, 1):
            sheet.column_dimensions[get_column_letter(index)].width = width

    @staticmethod
    def _header(sheet, values: Sequence[str]) -> None:
        cells = []
        for value in values:
            cell = WriteOnlyCell(sheet, value=value)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cells.append(cell)
        sheet.append(cells)

    def _validate(self, path: Path, stats: dict[str, object],
                  exported: dict[str, object]) -> None:
        checks = (
            exported["versions"] == stats["versions"],
            exported["changes"] == stats["changes"],
            exported["ADD"] == stats["ADD"], exported["MOD"] == stats["MOD"],
            exported["DEL"] == stats["DEL"],
            sum(int(item["count"]) for item in exported["sheets"]) == exported["changes"],
            all(int(item["count"]) <= self.change_rows_per_sheet for item in exported["sheets"]),
        )
        if not all(checks) or not zipfile.is_zipfile(path):
            raise ReportValidationError("Contagens ou estrutura do relatório divergentes")
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            expected = ["Resumo", "Execuções", "Versões processadas"] + [
                str(item["name"]) for item in exported["sheets"]
            ] + ["Erros", "Integridade"]
            if workbook.sheetnames != expected:
                raise ReportValidationError("Worksheets esperadas não estão presentes")
            previous_id: int | None = None
            for item in exported["sheets"]:
                sheet = workbook[str(item["name"])]
                rows = sheet.iter_rows(min_row=2, values_only=True)
                count = 0
                first = last = None
                for row in rows:
                    change_id = row[0]
                    first = change_id if first is None else first
                    last = change_id
                    count += 1
                    if previous_id is not None and change_id <= previous_id:
                        raise ReportValidationError("IDs duplicados ou ordem global inválida no XLSX")
                    previous_id = change_id
                if count != int(item["count"]) or count + 1 > EXCEL_MAX_ROWS:
                    raise ReportValidationError("Quantidade de linhas da worksheet inválida")
                if first != item["first_id"] or last != item["last_id"]:
                    raise ReportValidationError("Faixa de IDs da worksheet divergente")
        finally:
            workbook.close()

    def _progress(self, message: str, *, force: bool = False) -> None:
        now = time.monotonic()
        if self.progress_callback is not None and (force or now - self._last_progress >= 0.25):
            self.progress_callback(message)
            self._last_progress = now
