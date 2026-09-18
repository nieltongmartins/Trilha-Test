import os
import re
import json
import time
import queue
import threading
import difflib
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.formula.tokenizer import Tokenizer


APP_TITLE = "Comparador de Planilhas Excel"

SUPPORTED_EXTENSIONS = [
    ("Planilhas Excel", "*.xlsx *.xlsm"),
    ("Todos os arquivos", "*.*"),
]

# Quantidade máxima de itens usados para formar o perfil de uma linha/coluna.
# Mantém a detecção estrutural rápida mesmo em planilhas grandes.
PROFILE_SAMPLE_LIMIT = 40

# Quantos itens à frente serão examinados quando houver suspeita de
# linha/coluna inserida ou removida.
ALIGN_LOOKAHEAD = 10

# Similaridade mínima para considerar dois perfis equivalentes.
ALIGN_THRESHOLD = 0.60


# ============================================================
# FUNÇÕES BÁSICAS
# ============================================================

def has_content(value):
    if value is None:
        return False

    if isinstance(value, str) and value.strip() == "":
        return False

    return True


def normalize_text(value):
    if value is None:
        return ""

    text = str(value).strip().lower()
    text = re.sub(r"\s+", " ", text)

    return text


def normalize_profile_value(value):
    """
    Normalização usada somente para detectar estrutura.

    Fórmulas são tratadas como um tipo de conteúdo, em vez de comparar
    a fórmula exata, porque uma inserção de linha/coluna pode alterar
    automaticamente suas referências.
    """
    if value is None:
        return ""

    if isinstance(value, str):
        text = value.strip()

        if text.startswith("="):
            return "<FORMULA>"

        return "T:" + normalize_text(text)[:80]

    if isinstance(value, bool):
        return "B:" + str(value)

    if isinstance(value, (int, float)):
        return "N:" + str(value)

    if isinstance(value, datetime):
        return "D:" + value.isoformat()

    return "O:" + normalize_text(value)[:80]


def color_to_text(color):
    if color is None:
        return None

    data = {}

    for attr in ("type", "rgb", "indexed", "theme", "tint", "auto"):
        try:
            value = getattr(color, attr, None)

            if value is not None:
                data[attr] = value
        except Exception:
            pass

    return data or None


def side_to_dict(side):
    if side is None:
        return None

    return {
        "style": side.style,
        "color": color_to_text(side.color),
    }


def style_part(cell, part):
    if part == "font":
        font = cell.font

        return {
            "name": font.name,
            "size": font.sz,
            "bold": font.b,
            "italic": font.i,
            "underline": font.u,
            "strike": font.strike,
            "color": color_to_text(font.color),
            "vertAlign": font.vertAlign,
        }

    if part == "fill":
        fill = cell.fill

        return {
            "fill_type": fill.fill_type,
            "fgColor": color_to_text(fill.fgColor),
            "bgColor": color_to_text(fill.bgColor),
        }

    if part == "border":
        border = cell.border

        return {
            "left": side_to_dict(border.left),
            "right": side_to_dict(border.right),
            "top": side_to_dict(border.top),
            "bottom": side_to_dict(border.bottom),
            "diagonal": side_to_dict(border.diagonal),
            "vertical": side_to_dict(border.vertical),
            "horizontal": side_to_dict(border.horizontal),
            "diagonalUp": border.diagonalUp,
            "diagonalDown": border.diagonalDown,
            "outline": border.outline,
        }

    if part == "alignment":
        alignment = cell.alignment

        return {
            "horizontal": alignment.horizontal,
            "vertical": alignment.vertical,
            "text_rotation": alignment.textRotation,
            "wrap_text": alignment.wrapText,
            "shrink_to_fit": alignment.shrinkToFit,
            "indent": alignment.indent,
            "relative_indent": alignment.relativeIndent,
            "justify_last_line": alignment.justifyLastLine,
            "reading_order": alignment.readingOrder,
        }

    if part == "number_format":
        return cell.number_format

    if part == "protection":
        protection = cell.protection

        return {
            "locked": protection.locked,
            "hidden": protection.hidden,
        }

    return None


def safe_json(value):
    if value is None:
        return ""

    return json.dumps(
        value,
        ensure_ascii=False,
        default=str,
        sort_keys=True,
    )


def add_difference(
    differences,
    sheet,
    location,
    category,
    base,
    evaluated,
    details,
):
    differences.append({
        "sheet": sheet,
        "location": location,
        "category": category,
        "base": base,
        "evaluated": evaluated,
        "details": details,
    })


# ============================================================
# CÉLULAS / ÁREA UTILIZADA
# ============================================================

def get_meaningful_coordinates(ws):
    """
    Retorna somente células que realmente possuem conteúdo.

    Evita percorrer milhares/milhões de células vazias geradas por
    formatação antiga ou max_row/max_column exagerados.
    """
    coords = set()

    for key, cell in ws._cells.items():
        if not isinstance(key, tuple) or len(key) != 2:
            continue

        if has_content(getattr(cell, "value", None)):
            coords.add((cell.row, cell.column))

    return coords


def get_styled_coordinates(ws):
    """
    Retorna células materializadas no arquivo que possuem estilo,
    mesmo quando não possuem conteúdo.

    Assim alterações de borda, fonte, preenchimento, alinhamento,
    formato numérico ou proteção também são detectadas em células vazias.
    """
    coords = set()

    for key, cell in ws._cells.items():
        if not isinstance(key, tuple) or len(key) != 2:
            continue

        try:
            if cell.has_style:
                coords.add((cell.row, cell.column))
        except Exception:
            pass

    return coords


def get_relevant_coordinates(ws, include_styles=False):
    coords = get_meaningful_coordinates(ws)

    if include_styles:
        coords |= get_styled_coordinates(ws)

    return coords


def get_used_rows(coords):
    return sorted({row for row, _ in coords})


def get_used_columns(coords):
    return sorted({col for _, col in coords})


# ============================================================
# PERFIS PARA DETECÇÃO DE LINHAS/COLUNAS DESLOCADAS
# ============================================================

def make_column_profile(ws, col, used_rows):
    tokens = []

    for row in used_rows:
        value = ws.cell(row=row, column=col).value

        if has_content(value):
            tokens.append(normalize_profile_value(value))

            if len(tokens) >= PROFILE_SAMPLE_LIMIT:
                break

    if not tokens:
        return ("<EMPTY>",)

    return tuple(tokens)


def make_row_profile(ws, row, used_columns, column_map=None, is_base=False):
    """
    Para linhas, usamos somente colunas que possuem correspondência estrutural.

    Isso evita que uma coluna inserida faça todas as linhas parecerem diferentes.
    """
    tokens = []

    if column_map and is_base:
        columns = [
            col
            for col in used_columns
            if col in column_map
        ]
    else:
        columns = used_columns

    for col in columns:
        value = ws.cell(row=row, column=col).value

        if has_content(value):
            tokens.append(normalize_profile_value(value))

            if len(tokens) >= PROFILE_SAMPLE_LIMIT:
                break

    if not tokens:
        return ("<EMPTY>",)

    return tuple(tokens)


def profile_similarity(profile_a, profile_b):
    if profile_a == profile_b:
        return 1.0

    if not profile_a or not profile_b:
        return 0.0

    return difflib.SequenceMatcher(
        None,
        profile_a,
        profile_b,
        autojunk=False,
    ).ratio()


def _identity_mapping(base_items, eval_items):
    common = sorted(set(base_items) & set(eval_items))
    return {item: item for item in common}


def _has_real_structural_shift(mapping):
    """
    Só aceita hipótese estrutural quando existem âncoras suficientes
    demonstrando deslocamento consistente.
    """
    if not mapping or len(mapping) < 2:
        return False

    pairs = sorted(mapping.items())
    deltas = [eval_item - base_item for base_item, eval_item in pairs]

    if all(delta == 0 for delta in deltas):
        return False

    # Procura ao menos duas âncoras com o mesmo deslocamento diferente de zero.
    counts = {}

    for delta in deltas:
        if delta != 0:
            counts[delta] = counts.get(delta, 0) + 1

    return any(count >= 2 for count in counts.values())


def greedy_align(
    base_items,
    eval_items,
    base_profiles,
    eval_profiles,
    threshold=ALIGN_THRESHOLD,
    lookahead=ALIGN_LOOKAHEAD,
):
    """
    Alinhamento estrutural conservador.

    Regra central:
    conteúdo existir só em um dos arquivos NÃO significa que houve
    inserção/exclusão de linha ou coluna.

    A hipótese estrutural só é mantida quando existem âncoras reais que
    comprovam um deslocamento consistente.
    """
    if not base_items or not eval_items:
        return _identity_mapping(base_items, eval_items), [], []

    mapping = {}
    removed = []
    added = []

    i = 0
    j = 0

    while i < len(base_items) and j < len(eval_items):
        base_item = base_items[i]
        eval_item = eval_items[j]

        current_similarity = profile_similarity(
            base_profiles[base_item],
            eval_profiles[eval_item],
        )

        if current_similarity >= threshold:
            mapping[base_item] = eval_item
            i += 1
            j += 1
            continue

        best_eval_offset = None
        best_eval_score = 0.0

        max_eval_offset = min(
            lookahead,
            len(eval_items) - j - 1,
        )

        for offset in range(1, max_eval_offset + 1):
            candidate = eval_items[j + offset]

            score = profile_similarity(
                base_profiles[base_item],
                eval_profiles[candidate],
            )

            adjusted_score = score - (offset * 0.01)

            if adjusted_score > best_eval_score:
                best_eval_score = adjusted_score
                best_eval_offset = offset

        best_base_offset = None
        best_base_score = 0.0

        max_base_offset = min(
            lookahead,
            len(base_items) - i - 1,
        )

        for offset in range(1, max_base_offset + 1):
            candidate = base_items[i + offset]

            score = profile_similarity(
                base_profiles[candidate],
                eval_profiles[eval_item],
            )

            adjusted_score = score - (offset * 0.01)

            if adjusted_score > best_base_score:
                best_base_score = adjusted_score
                best_base_offset = offset

        found_eval_match = (
            best_eval_offset is not None
            and best_eval_score >= threshold - 0.02
        )

        found_base_match = (
            best_base_offset is not None
            and best_base_score >= threshold - 0.02
        )

        if found_eval_match and (
            not found_base_match
            or best_eval_score >= best_base_score
        ):
            for offset in range(best_eval_offset):
                added.append(eval_items[j + offset])

            j += best_eval_offset
            continue

        if found_base_match:
            for offset in range(best_base_offset):
                removed.append(base_items[i + offset])

            i += best_base_offset
            continue

        # Sem evidência de deslocamento, compara pela posição física.
        mapping[base_item] = eval_item
        i += 1
        j += 1

    while i < len(base_items):
        removed.append(base_items[i])
        i += 1

    while j < len(eval_items):
        added.append(eval_items[j])
        j += 1

    # Sem deslocamento comprovado, abandona a hipótese estrutural.
    if not _has_real_structural_shift(mapping):
        return _identity_mapping(base_items, eval_items), [], []

    return mapping, removed, added


def build_comparison_axes(
    coords_base,
    coords_eval,
    row_map,
    column_map,
    removed_rows=None,
    added_rows=None,
    removed_columns=None,
    added_columns=None,
):
    """
    Completa o mapeamento para conteúdo/estilo presente somente em um arquivo,
    mas NÃO recria posições que já foram comprovadas como inserção/remoção
    estrutural.

    Assim uma linha/coluna realmente inserida aparece uma vez como evento
    estrutural, em vez de também gerar dezenas de diferenças célula a célula.
    """
    removed_rows = set(removed_rows or [])
    added_rows = set(added_rows or [])
    removed_columns = set(removed_columns or [])
    added_columns = set(added_columns or [])

    base_rows = get_used_rows(coords_base)
    eval_rows = get_used_rows(coords_eval)

    base_cols = get_used_columns(coords_base)
    eval_cols = get_used_columns(coords_eval)

    mapped_rows = dict(row_map)

    mapped_base_rows = set(mapped_rows.keys())
    mapped_eval_rows = set(mapped_rows.values())

    for row in sorted(set(base_rows) | set(eval_rows)):
        if row in removed_rows or row in added_rows:
            continue

        if row not in mapped_base_rows and row not in mapped_eval_rows:
            mapped_rows[row] = row
            mapped_base_rows.add(row)
            mapped_eval_rows.add(row)

    mapped_cols = dict(column_map)

    mapped_base_cols = set(mapped_cols.keys())
    mapped_eval_cols = set(mapped_cols.values())

    for col in sorted(set(base_cols) | set(eval_cols)):
        if col in removed_columns or col in added_columns:
            continue

        if col not in mapped_base_cols and col not in mapped_eval_cols:
            mapped_cols[col] = col
            mapped_base_cols.add(col)
            mapped_eval_cols.add(col)

    return mapped_rows, mapped_cols


def derive_structural_gaps(mapping):
    """
    Deriva inserções/exclusões físicas usando pares de âncoras correspondentes.

    Exemplo:
        Base:     linha 8 -> linha 9
        Avaliada: linha 8 -> linha 10

    Existe uma linha física adicional entre as duas âncoras no Avaliado:
        linha 9 adicionada

    O mesmo vale para colunas.

    Como a decisão exige âncoras em ambos os lados do ponto alterado,
    simplesmente preencher uma célula antes vazia NÃO é confundido com
    inserção estrutural.
    """
    if not mapping or len(mapping) < 2:
        return [], []

    pairs = sorted(mapping.items())
    removed = []
    added = []

    for (base_a, eval_a), (base_b, eval_b) in zip(pairs, pairs[1:]):
        if base_b <= base_a or eval_b <= eval_a:
            continue

        base_gap = base_b - base_a - 1
        eval_gap = eval_b - eval_a - 1

        if eval_gap > base_gap:
            difference = eval_gap - base_gap

            # As posições extras ficam imediatamente depois da âncora anterior.
            for position in range(
                eval_a + 1,
                eval_a + 1 + difference,
            ):
                added.append(position)

        elif base_gap > eval_gap:
            difference = base_gap - eval_gap

            for position in range(
                base_a + 1,
                base_a + 1 + difference,
            ):
                removed.append(position)

    # Inserção/remoção no início pode ser reconhecida quando há ao menos
    # duas âncoras com o mesmo deslocamento, o que dá evidência suficiente.
    deltas = [
        eval_pos - base_pos
        for base_pos, eval_pos in pairs
    ]

    if len(pairs) >= 2 and len(set(deltas[:2])) == 1:
        first_delta = deltas[0]

        if first_delta > 0:
            for position in range(1, first_delta + 1):
                if position not in added:
                    added.append(position)

        elif first_delta < 0:
            for position in range(1, abs(first_delta) + 1):
                if position not in removed:
                    removed.append(position)

    return sorted(set(removed)), sorted(set(added))


def worksheet_protection_snapshot(ws):
    protection = ws.protection

    attrs = [
        "sheet",
        "objects",
        "scenarios",
        "formatCells",
        "formatColumns",
        "formatRows",
        "insertColumns",
        "insertRows",
        "insertHyperlinks",
        "deleteColumns",
        "deleteRows",
        "selectLockedCells",
        "selectUnlockedCells",
        "sort",
        "autoFilter",
        "pivotTables",
    ]

    result = {}

    for attr in attrs:
        try:
            result[attr] = getattr(protection, attr)
        except Exception:
            result[attr] = None

    return result

def detect_column_structure(ws_base, ws_eval, coords_base, coords_eval):
    base_columns = get_used_columns(coords_base)
    eval_columns = get_used_columns(coords_eval)

    base_rows = get_used_rows(coords_base)
    eval_rows = get_used_rows(coords_eval)

    base_profiles = {
        col: make_column_profile(
            ws_base,
            col,
            base_rows,
        )
        for col in base_columns
    }

    eval_profiles = {
        col: make_column_profile(
            ws_eval,
            col,
            eval_rows,
        )
        for col in eval_columns
    }

    mapping, _, _ = greedy_align(
        base_columns,
        eval_columns,
        base_profiles,
        eval_profiles,
    )

    removed, added = derive_structural_gaps(mapping)

    return mapping, removed, added


def detect_row_structure(
    ws_base,
    ws_eval,
    coords_base,
    coords_eval,
    column_map,
):
    base_rows = get_used_rows(coords_base)
    eval_rows = get_used_rows(coords_eval)

    base_columns = get_used_columns(coords_base)

    # Para o Avaliado, usa somente as colunas que correspondem às colunas Base.
    eval_columns = sorted(set(column_map.values()))

    base_profiles = {
        row: make_row_profile(
            ws_base,
            row,
            base_columns,
            column_map=column_map,
            is_base=True,
        )
        for row in base_rows
    }

    eval_profiles = {
        row: make_row_profile(
            ws_eval,
            row,
            eval_columns,
        )
        for row in eval_rows
    }

    mapping, _, _ = greedy_align(
        base_rows,
        eval_rows,
        base_profiles,
        eval_profiles,
    )

    removed, added = derive_structural_gaps(mapping)

    return mapping, removed, added



def get_merged_range_for_cell(ws, row, col):
    """
    Retorna o intervalo mesclado que contém a célula, se existir.
    """
    for merged in ws.merged_cells.ranges:
        if (
            merged.min_row <= row <= merged.max_row
            and merged.min_col <= col <= merged.max_col
        ):
            return merged

    return None


def get_effective_cell_value(ws, row, col):
    """
    Em células mescladas, o openpyxl armazena o valor somente na
    célula superior esquerda. Esta função retorna esse valor.
    """
    merged = get_merged_range_for_cell(
        ws,
        row,
        col,
    )

    if merged is not None:
        anchor = ws.cell(
            row=merged.min_row,
            column=merged.min_col,
        )

        return anchor.value, merged

    cell = ws.cell(
        row=row,
        column=col,
    )

    return cell.value, None


def report_structural_area_content(
    differences,
    sheet_name,
    ws_base,
    ws_eval,
    added_rows,
    removed_rows,
    added_columns,
    removed_columns,
    options,
):
    """
    Registra conteúdo existente dentro de linhas/colunas que foram
    estruturalmente adicionadas ou removidas.

    Antes, essas áreas eram excluídas da comparação para evitar cascata.
    Agora o evento estrutural continua sendo reportado, mas o conteúdo
    real dentro da área também aparece no relatório.

    Células mescladas são tratadas pelo intervalo completo e pelo valor
    armazenado na célula âncora superior esquerda.
    """
    emitted_added = set()
    emitted_removed = set()

    # --------------------------------------------------------
    # CONTEÚDO ADICIONADO EM LINHAS NOVAS
    # --------------------------------------------------------
    if options.get("added_content", False):
        for row in sorted(set(added_rows or [])):
            for key, cell in ws_eval._cells.items():
                if not isinstance(key, tuple) or len(key) != 2:
                    continue

                if cell.row != row:
                    continue

                value, merged = get_effective_cell_value(
                    ws_eval,
                    cell.row,
                    cell.column,
                )

                if not has_content(value):
                    continue

                if merged is not None:
                    identity = (
                        merged.min_row,
                        merged.min_col,
                        merged.max_row,
                        merged.max_col,
                    )

                    if identity in emitted_added:
                        continue

                    emitted_added.add(identity)

                    location = (
                        f"{sheet_name}!{str(merged)}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma linha adicionada. "
                        "A célula pertence a um intervalo mesclado; o valor "
                        "está armazenado na célula superior esquerda da mesclagem."
                    )
                else:
                    identity = (
                        cell.row,
                        cell.column,
                    )

                    if identity in emitted_added:
                        continue

                    emitted_added.add(identity)

                    location = (
                        f"{sheet_name}!"
                        f"{get_column_letter(cell.column)}{cell.row}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma linha adicionada."
                    )

                add_difference(
                    differences,
                    sheet_name,
                    location,
                    "Conteúdo adicionado",
                    "",
                    value,
                    details,
                )

        # ----------------------------------------------------
        # CONTEÚDO ADICIONADO EM COLUNAS NOVAS
        # ----------------------------------------------------
        for col in sorted(set(added_columns or [])):
            for key, cell in ws_eval._cells.items():
                if not isinstance(key, tuple) or len(key) != 2:
                    continue

                if cell.column != col:
                    continue

                value, merged = get_effective_cell_value(
                    ws_eval,
                    cell.row,
                    cell.column,
                )

                if not has_content(value):
                    continue

                if merged is not None:
                    identity = (
                        merged.min_row,
                        merged.min_col,
                        merged.max_row,
                        merged.max_col,
                    )

                    if identity in emitted_added:
                        continue

                    emitted_added.add(identity)

                    location = (
                        f"{sheet_name}!{str(merged)}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma coluna adicionada. "
                        "A célula pertence a um intervalo mesclado."
                    )
                else:
                    identity = (
                        cell.row,
                        cell.column,
                    )

                    if identity in emitted_added:
                        continue

                    emitted_added.add(identity)

                    location = (
                        f"{sheet_name}!"
                        f"{get_column_letter(cell.column)}{cell.row}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma coluna adicionada."
                    )

                add_difference(
                    differences,
                    sheet_name,
                    location,
                    "Conteúdo adicionado",
                    "",
                    value,
                    details,
                )

    # --------------------------------------------------------
    # CONTEÚDO REMOVIDO EM LINHAS/COLUNAS EXCLUÍDAS
    # --------------------------------------------------------
    if options.get("removed_content", False):
        for row in sorted(set(removed_rows or [])):
            for key, cell in ws_base._cells.items():
                if not isinstance(key, tuple) or len(key) != 2:
                    continue

                if cell.row != row:
                    continue

                value, merged = get_effective_cell_value(
                    ws_base,
                    cell.row,
                    cell.column,
                )

                if not has_content(value):
                    continue

                if merged is not None:
                    identity = (
                        merged.min_row,
                        merged.min_col,
                        merged.max_row,
                        merged.max_col,
                    )

                    if identity in emitted_removed:
                        continue

                    emitted_removed.add(identity)

                    location = (
                        f"{sheet_name}!{str(merged)}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma linha removida. "
                        "A célula pertencia a um intervalo mesclado."
                    )
                else:
                    identity = (
                        cell.row,
                        cell.column,
                    )

                    if identity in emitted_removed:
                        continue

                    emitted_removed.add(identity)

                    location = (
                        f"{sheet_name}!"
                        f"{get_column_letter(cell.column)}{cell.row}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma linha removida."
                    )

                add_difference(
                    differences,
                    sheet_name,
                    location,
                    "Conteúdo removido",
                    value,
                    "",
                    details,
                )

        for col in sorted(set(removed_columns or [])):
            for key, cell in ws_base._cells.items():
                if not isinstance(key, tuple) or len(key) != 2:
                    continue

                if cell.column != col:
                    continue

                value, merged = get_effective_cell_value(
                    ws_base,
                    cell.row,
                    cell.column,
                )

                if not has_content(value):
                    continue

                if merged is not None:
                    identity = (
                        merged.min_row,
                        merged.min_col,
                        merged.max_row,
                        merged.max_col,
                    )

                    if identity in emitted_removed:
                        continue

                    emitted_removed.add(identity)

                    location = (
                        f"{sheet_name}!{str(merged)}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma coluna removida. "
                        "A célula pertencia a um intervalo mesclado."
                    )
                else:
                    identity = (
                        cell.row,
                        cell.column,
                    )

                    if identity in emitted_removed:
                        continue

                    emitted_removed.add(identity)

                    location = (
                        f"{sheet_name}!"
                        f"{get_column_letter(cell.column)}{cell.row}"
                    )

                    details = (
                        "Conteúdo existente dentro de uma coluna removida."
                    )

                add_difference(
                    differences,
                    sheet_name,
                    location,
                    "Conteúdo removido",
                    value,
                    "",
                    details,
                )

# ============================================================
# FÓRMULAS COM REFERÊNCIAS REPOSICIONADAS
# ============================================================

CELL_REF_RE = re.compile(
    r"(?<![A-Z0-9_])(\$?)([A-Z]{1,3})(\$?)(\d+)",
    re.IGNORECASE,
)


def infer_mapped_position(position, mapping):
    """
    Retorna a posição equivalente no arquivo Avaliado.

    Prioridade:
    1. mapeamento estrutural exato;
    2. deslocamento inferido pelos vizinhos estruturais.

    A inferência só é usada quando existe evidência consistente.
    Caso contrário, mantém a posição original, evitando inventar deslocamentos.
    """
    if position in mapping:
        return mapping[position]

    if not mapping:
        return position

    keys = sorted(mapping)

    previous_keys = [
        key for key in keys
        if key < position
    ]

    next_keys = [
        key for key in keys
        if key > position
    ]

    previous_delta = None
    next_delta = None

    if previous_keys:
        key = previous_keys[-1]
        previous_delta = mapping[key] - key

    if next_keys:
        key = next_keys[0]
        next_delta = mapping[key] - key

    # Se os dois lados concordam, o deslocamento é seguro.
    if (
        previous_delta is not None
        and next_delta is not None
        and previous_delta == next_delta
    ):
        return position + previous_delta

    # Se só existe evidência de um lado, utiliza o deslocamento desse lado.
    if previous_delta is not None and next_delta is None:
        return position + previous_delta

    if next_delta is not None and previous_delta is None:
        return position + next_delta

    # Em uma fronteira de inserção/remoção os deltas podem ser diferentes.
    # Nesse caso, não inventamos equivalência.
    return position


def translate_reference_token(token_value, row_map, column_map):
    """
    Traduz referências A1 existentes em um token do tipo RANGE.

    Exemplos:
        $G$4:G4  -> $I$4:I4
        H4       -> J4
        O4       -> Q4

    Apenas referências reconhecidas como referências Excel são alteradas.
    Textos, operadores, nomes de funções e constantes permanecem intactos.
    """
    def replace_ref(match):
        col_abs = match.group(1)
        col_letters = match.group(2).upper()
        row_abs = match.group(3)
        row_number = int(match.group(4))

        try:
            col_number = column_index_from_string(
                col_letters
            )
        except Exception:
            return match.group(0)

        mapped_col = infer_mapped_position(
            col_number,
            column_map,
        )

        mapped_row = infer_mapped_position(
            row_number,
            row_map,
        )

        return (
            f"{col_abs}"
            f"{get_column_letter(mapped_col)}"
            f"{row_abs}"
            f"{mapped_row}"
        )

    return CELL_REF_RE.sub(
        replace_ref,
        token_value,
    )


def tokenize_formula_for_comparison(
    formula,
    row_map=None,
    column_map=None,
    translate=False,
):
    """
    Converte a fórmula em sequência de tokens comparáveis.

    Segurança contra falso positivo:
    - somente tokens classificados pelo openpyxl como RANGE têm suas
      referências reposicionadas;
    - strings como "Reprovado" ou "Cancelado" nunca são tratadas como célula;
    - funções, operadores, números e textos precisam continuar idênticos;
    - se qualquer parte lógica mudar, a fórmula continua sendo reportada.
    """
    if not isinstance(formula, str):
        return None

    if not formula.startswith("="):
        return None

    try:
        tokenizer = Tokenizer(formula)
    except Exception:
        return None

    result = []

    for token in tokenizer.items:
        value = token.value

        if (
            translate
            and token.type == "OPERAND"
            and token.subtype == "RANGE"
        ):
            value = translate_reference_token(
                value,
                row_map or {},
                column_map or {},
            )

        result.append(
            (
                token.type,
                token.subtype,
                value,
            )
        )

    return result


def remap_formula_references(
    formula,
    row_map,
    column_map,
):
    """
    Retorna uma versão textual aproximada da fórmula Base reposicionada.

    É usada principalmente para explicar o resultado no relatório.
    A decisão de equivalência é feita pela comparação token a token.
    """
    tokens = tokenize_formula_for_comparison(
        formula,
        row_map=row_map,
        column_map=column_map,
        translate=True,
    )

    if tokens is None:
        return formula

    return "".join(
        token_value
        for _, _, token_value in tokens
    )


def formulas_equivalent(
    base_formula,
    eval_formula,
    row_map,
    column_map,
):
    """
    Considera duas fórmulas equivalentes somente quando:

    1. a estrutura lógica/tokenizada é igual;
    2. toda diferença de referência é explicada pelo mapeamento
       estrutural de linhas/colunas;
    3. nenhuma constante, texto, operador ou função foi alterado.

    Dessa forma:
        =SE(H4="";M4;...)
    pode equivaler a
        =SE(J4="";O4;...)
    quando o sistema detectou duas colunas inseridas antes dessas referências.

    Porém:
        =SE(H4="";M4;0)
    NÃO será equivalente a
        =SE(J4="";O4;1)

    porque a constante mudou de 0 para 1.
    """
    if base_formula == eval_formula:
        return True

    base_tokens = tokenize_formula_for_comparison(
        base_formula,
        row_map=row_map,
        column_map=column_map,
        translate=True,
    )

    eval_tokens = tokenize_formula_for_comparison(
        eval_formula,
        translate=False,
    )

    if (
        base_tokens is None
        or eval_tokens is None
    ):
        return False

    return base_tokens == eval_tokens


# ============================================================
# MESCLAGENS
# ============================================================

def merged_ranges_relevant(ws, meaningful_coords):
    relevant = set()

    if not meaningful_coords:
        return relevant

    for merged in ws.merged_cells.ranges:
        for row, col in meaningful_coords:
            if (
                merged.min_row <= row <= merged.max_row
                and merged.min_col <= col <= merged.max_col
            ):
                relevant.add(str(merged))
                break

    return relevant


# ============================================================
# PROGRESSO
# ============================================================

class ProgressTracker:
    def __init__(self, callback, total_units):
        self.callback = callback
        self.total_units = max(total_units, 1)
        self.completed = 0

    def update(self, units=1, message=""):
        self.completed += units

        percent = min(
            100.0,
            (self.completed / self.total_units) * 100.0,
        )

        if self.callback:
            self.callback(
                percent,
                message,
            )

    def set_message(self, message):
        if self.callback:
            percent = min(
                100.0,
                (self.completed / self.total_units) * 100.0,
            )

            self.callback(
                percent,
                message,
            )


# ============================================================
# COMPARAÇÃO PRINCIPAL
# ============================================================

def compare_workbooks(
    base_path,
    evaluated_path,
    options,
    progress_callback=None,
):
    if progress_callback:
        progress_callback(
            1.0,
            "Abrindo as planilhas...",
        )

    wb_base = load_workbook(
        base_path,
        data_only=False,
    )

    wb_eval = load_workbook(
        evaluated_path,
        data_only=False,
    )

    differences = []

    base_sheets = set(wb_base.sheetnames)
    eval_sheets = set(wb_eval.sheetnames)

    if options["sheets"]:
        for name in sorted(base_sheets - eval_sheets):
            add_difference(
                differences,
                name,
                name,
                "Planilha ausente",
                "Existe",
                "Não existe",
                "A planilha existe no Base, mas não no Avaliado.",
            )

        for name in sorted(eval_sheets - base_sheets):
            add_difference(
                differences,
                name,
                name,
                "Planilha adicional",
                "Não existe",
                "Existe",
                "A planilha existe no Avaliado, mas não no Base.",
            )

    common_sheets = sorted(
        base_sheets & eval_sheets
    )

    # Pré-cálculo leve para estimar o trabalho total.
    sheet_data = {}

    estimated_units = 0

    for sheet_name in common_sheets:
        ws_base = wb_base[sheet_name]
        ws_eval = wb_eval[sheet_name]

        coords_base = get_meaningful_coordinates(
            ws_base
        )

        coords_eval = get_meaningful_coordinates(
            ws_eval
        )

        sheet_data[sheet_name] = (
            coords_base,
            coords_eval,
        )

        estimated_units += (
            len(coords_base)
            + len(coords_eval)
            + 20
        )

    tracker = ProgressTracker(
        progress_callback,
        estimated_units,
    )

    selected_style_parts = []

    if options["font"]:
        selected_style_parts.append(
            ("font", "Fonte")
        )

    if options["fill"]:
        selected_style_parts.append(
            ("fill", "Preenchimento")
        )

    if options["border"]:
        selected_style_parts.append(
            ("border", "Borda")
        )

    if options["alignment"]:
        selected_style_parts.append(
            ("alignment", "Alinhamento")
        )

    if options["number_format"]:
        selected_style_parts.append(
            ("number_format", "Formato numérico")
        )

    if options["protection"]:
        selected_style_parts.append(
            ("protection", "Proteção da célula")
        )

    include_style_cells = bool(selected_style_parts)

    for sheet_index, sheet_name in enumerate(
        common_sheets,
        start=1,
    ):
        ws_base = wb_base[sheet_name]
        ws_eval = wb_eval[sheet_name]

        coords_base, coords_eval = (
            sheet_data[sheet_name]
        )

        compare_coords_base = get_relevant_coordinates(
            ws_base,
            include_styles=include_style_cells,
        )

        compare_coords_eval = get_relevant_coordinates(
            ws_eval,
            include_styles=include_style_cells,
        )

        tracker.set_message(
            f"Aba {sheet_index}/{len(common_sheets)} - "
            f"detectando deslocamentos estruturais..."
        )

        # --------------------------------------------------------
        # 1. COLUNAS
        # --------------------------------------------------------
        column_map, removed_columns, added_columns = (
            detect_column_structure(
                ws_base,
                ws_eval,
                coords_base,
                coords_eval,
            )
        )

        if options["removed_columns"]:
            for col in removed_columns:
                letter = get_column_letter(col)

                add_difference(
                    differences,
                    sheet_name,
                    f"{sheet_name}!Coluna {letter}",
                    "Coluna removida",
                    letter,
                    "",
                    (
                        f"A coluna {letter} existente no Base não encontrou "
                        f"correspondência estrutural no Avaliado."
                    ),
                )

        if options["added_columns"]:
            for col in added_columns:
                letter = get_column_letter(col)

                add_difference(
                    differences,
                    sheet_name,
                    f"{sheet_name}!Coluna {letter}",
                    "Coluna adicionada",
                    "",
                    letter,
                    (
                        f"A coluna {letter} foi identificada como nova no "
                        f"arquivo Avaliado."
                    ),
                )

        # --------------------------------------------------------
        # 2. LINHAS
        # --------------------------------------------------------
        row_map, removed_rows, added_rows = (
            detect_row_structure(
                ws_base,
                ws_eval,
                coords_base,
                coords_eval,
                column_map,
            )
        )

        if options["removed_rows"]:
            for row in removed_rows:
                add_difference(
                    differences,
                    sheet_name,
                    f"{sheet_name}!Linha {row}",
                    "Linha removida",
                    row,
                    "",
                    (
                        f"A linha {row} existente no Base não encontrou "
                        f"correspondência estrutural no Avaliado."
                    ),
                )

        if options["added_rows"]:
            for row in added_rows:
                add_difference(
                    differences,
                    sheet_name,
                    f"{sheet_name}!Linha {row}",
                    "Linha adicionada",
                    "",
                    row,
                    (
                        f"A linha {row} foi identificada como nova no "
                        f"arquivo Avaliado."
                    ),
                )

        # Mesmo quando há alteração estrutural, o conteúdo existente
        # dentro da linha/coluna adicionada ou removida também deve
        # ser registrado, inclusive em células mescladas.
        report_structural_area_content(
            differences,
            sheet_name,
            ws_base,
            ws_eval,
            added_rows,
            removed_rows,
            added_columns,
            removed_columns,
            options,
        )

        tracker.update(
            10,
            f"Aba {sheet_index}/{len(common_sheets)} - "
            f"estrutura alinhada.",
        )

        # --------------------------------------------------------
        # 3. CONFIGURAÇÕES GERAIS DA ABA
        # --------------------------------------------------------
        if options["sheet_state"]:
            if ws_base.sheet_state != ws_eval.sheet_state:
                add_difference(
                    differences,
                    sheet_name,
                    sheet_name,
                    "Estado da aba",
                    ws_base.sheet_state,
                    ws_eval.sheet_state,
                    "O estado da aba (visível/oculta) está diferente.",
                )

        if options.get("sheet_protection", False):
            base_sheet_protection = worksheet_protection_snapshot(ws_base)
            eval_sheet_protection = worksheet_protection_snapshot(ws_eval)

            if base_sheet_protection != eval_sheet_protection:
                add_difference(
                    differences,
                    sheet_name,
                    sheet_name,
                    "Proteção da planilha",
                    safe_json(base_sheet_protection),
                    safe_json(eval_sheet_protection),
                    "As configurações de proteção da planilha estão diferentes.",
                )

        if options["freeze_panes"]:
            base_freeze = str(
                ws_base.freeze_panes or ""
            )

            eval_freeze = str(
                ws_eval.freeze_panes or ""
            )

            if base_freeze != eval_freeze:
                add_difference(
                    differences,
                    sheet_name,
                    sheet_name,
                    "Painel congelado",
                    base_freeze,
                    eval_freeze,
                    "A configuração de painel congelado está diferente.",
                )

        # --------------------------------------------------------
        # 4. COMPARAÇÃO CÉLULA A CÉLULA, MAS JÁ ALINHADA
        # --------------------------------------------------------
        row_map, column_map = build_comparison_axes(
            compare_coords_base,
            compare_coords_eval,
            row_map,
            column_map,
            removed_rows=removed_rows,
            added_rows=added_rows,
            removed_columns=removed_columns,
            added_columns=added_columns,
        )

        mapped_rows = sorted(row_map.items())
        mapped_columns = sorted(column_map.items())

        total_pairs = max(
            1,
            len(mapped_rows) * max(len(mapped_columns), 1),
        )

        processed_pairs = 0
        progress_interval = max(
            100,
            total_pairs // 150,
        )

        for base_row, eval_row in mapped_rows:
            for base_col, eval_col in mapped_columns:
                processed_pairs += 1

                c_base = ws_base.cell(
                    row=base_row,
                    column=base_col,
                )

                c_eval = ws_eval.cell(
                    row=eval_row,
                    column=eval_col,
                )

                b_has = has_content(
                    c_base.value
                )

                e_has = has_content(
                    c_eval.value
                )

                base_coordinate = (
                    f"{get_column_letter(base_col)}{base_row}"
                )

                eval_coordinate = (
                    f"{get_column_letter(eval_col)}{eval_row}"
                )

                if (
                    base_coordinate == eval_coordinate
                ):
                    location = (
                        f"{sheet_name}!{base_coordinate}"
                    )
                else:
                    location = (
                        f"{sheet_name}!{base_coordinate}"
                        f" → {eval_coordinate}"
                    )

                # FORMATAÇÃO: é avaliada mesmo quando a célula está vazia.
                for style_key, label in selected_style_parts:
                    base_style = style_part(
                        c_base,
                        style_key,
                    )

                    eval_style = style_part(
                        c_eval,
                        style_key,
                    )

                    if base_style != eval_style:
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            f"Formatação - {label}",
                            safe_json(base_style),
                            safe_json(eval_style),
                            (
                                f"A formatação de {label.lower()} "
                                f"está diferente."
                            ),
                        )

                # Se ambas estiverem vazias, a análise de conteúdo termina aqui.
                # A formatação já foi verificada acima.
                if not b_has and not e_has:
                    continue

                if b_has and not e_has:
                    if options["removed_content"]:
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Conteúdo removido",
                            c_base.value,
                            "",
                            (
                                "A célula correspondente possui conteúdo no "
                                "Base e está vazia no Avaliado."
                            ),
                        )

                    continue

                if e_has and not b_has:
                    if options["added_content"]:
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Conteúdo adicionado",
                            "",
                            c_eval.value,
                            (
                                "A célula correspondente está vazia no Base "
                                "e possui conteúdo no Avaliado."
                            ),
                        )

                    continue

                base_is_formula = (
                    isinstance(c_base.value, str)
                    and c_base.value.startswith("=")
                )

                eval_is_formula = (
                    isinstance(c_eval.value, str)
                    and c_eval.value.startswith("=")
                )

                if (
                    options["formula"]
                    and (
                        base_is_formula
                        or eval_is_formula
                    )
                ):
                    if (
                        not base_is_formula
                        or not eval_is_formula
                    ):
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Fórmula alterada",
                            c_base.value,
                            c_eval.value,
                            (
                                "Uma das células contém fórmula e a outra não."
                            ),
                        )

                    elif not formulas_equivalent(
                        c_base.value,
                        c_eval.value,
                        row_map,
                        column_map,
                    ):
                        expected_formula = remap_formula_references(
                            c_base.value,
                            row_map,
                            column_map,
                        )

                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Fórmula alterada",
                            c_base.value,
                            c_eval.value,
                            (
                                "A fórmula contém uma alteração que não pode ser "
                                "explicada somente pelos deslocamentos estruturais "
                                "detectados. "
                                f"Fórmula Base reposicionada esperada: "
                                f"{expected_formula}"
                            ),
                        )

                elif (
                    options["value"]
                    and c_base.value != c_eval.value
                ):
                    add_difference(
                        differences,
                        sheet_name,
                        location,
                        "Valor alterado",
                        c_base.value,
                        c_eval.value,
                        "O valor da célula correspondente está diferente.",
                    )

                if (
                    processed_pairs % progress_interval
                    == 0
                ):
                    tracker.update(
                        progress_interval,
                        (
                            f"Aba {sheet_index}/{len(common_sheets)} - "
                            f"comparando células..."
                        ),
                    )

        # Compensa o restante da estimativa referente às células.
        tracker.update(
            max(
                1,
                len(coords_base)
                + len(coords_eval)
                - min(
                    processed_pairs,
                    len(coords_base)
                    + len(coords_eval),
                ),
            ),
            (
                f"Aba {sheet_index}/{len(common_sheets)} - "
                f"verificando propriedades estruturais..."
            ),
        )

        # --------------------------------------------------------
        # 5. ALTURA / OCULTAÇÃO DAS LINHAS MAPEADAS
        # --------------------------------------------------------
        if (
            options["row_height"]
            or options["row_hidden"]
        ):
            for base_row, eval_row in mapped_rows:
                dim_base = (
                    ws_base.row_dimensions[base_row]
                )

                dim_eval = (
                    ws_eval.row_dimensions[eval_row]
                )

                location = (
                    f"{sheet_name}!Linha {base_row}"
                )

                if base_row != eval_row:
                    location += (
                        f" → Linha {eval_row}"
                    )

                if options["row_height"]:
                    if (
                        dim_base.height
                        != dim_eval.height
                    ):
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Linha - altura",
                            dim_base.height,
                            dim_eval.height,
                            (
                                "A altura da linha correspondente "
                                "está diferente."
                            ),
                        )

                if options["row_hidden"]:
                    base_hidden = bool(
                        dim_base.hidden
                    )

                    eval_hidden = bool(
                        dim_eval.hidden
                    )

                    if (
                        base_hidden
                        != eval_hidden
                    ):
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Linha - oculta",
                            base_hidden,
                            eval_hidden,
                            (
                                "O estado de ocultação da linha "
                                "correspondente está diferente."
                            ),
                        )

        # --------------------------------------------------------
        # 6. LARGURA / OCULTAÇÃO DAS COLUNAS MAPEADAS
        # --------------------------------------------------------
        if (
            options["column_width"]
            or options["column_hidden"]
        ):
            for base_col, eval_col in mapped_columns:
                base_letter = (
                    get_column_letter(base_col)
                )

                eval_letter = (
                    get_column_letter(eval_col)
                )

                dim_base = (
                    ws_base.column_dimensions[
                        base_letter
                    ]
                )

                dim_eval = (
                    ws_eval.column_dimensions[
                        eval_letter
                    ]
                )

                location = (
                    f"{sheet_name}!Coluna "
                    f"{base_letter}"
                )

                if base_col != eval_col:
                    location += (
                        f" → Coluna {eval_letter}"
                    )

                if options["column_width"]:
                    if (
                        dim_base.width
                        != dim_eval.width
                    ):
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Coluna - largura",
                            dim_base.width,
                            dim_eval.width,
                            (
                                "A largura da coluna correspondente "
                                "está diferente."
                            ),
                        )

                if options["column_hidden"]:
                    base_hidden = bool(
                        dim_base.hidden
                    )

                    eval_hidden = bool(
                        dim_eval.hidden
                    )

                    if (
                        base_hidden
                        != eval_hidden
                    ):
                        add_difference(
                            differences,
                            sheet_name,
                            location,
                            "Coluna - oculta",
                            base_hidden,
                            eval_hidden,
                            (
                                "O estado de ocultação da coluna "
                                "correspondente está diferente."
                            ),
                        )

        # --------------------------------------------------------
        # 7. MESCLAGENS
        # --------------------------------------------------------
        if options["merged_cells"]:
            merges_base = (
                merged_ranges_relevant(
                    ws_base,
                    coords_base,
                )
            )

            merges_eval = (
                merged_ranges_relevant(
                    ws_eval,
                    coords_eval,
                )
            )

            for merged in sorted(
                merges_base - merges_eval
            ):
                add_difference(
                    differences,
                    sheet_name,
                    f"{sheet_name}!{merged}",
                    "Mesclagem removida/alterada",
                    merged,
                    "",
                    (
                        "A mesclagem existe no Base, "
                        "mas não no Avaliado."
                    ),
                )

            for merged in sorted(
                merges_eval - merges_base
            ):
                add_difference(
                    differences,
                    sheet_name,
                    f"{sheet_name}!{merged}",
                    "Mesclagem adicionada/alterada",
                    "",
                    merged,
                    (
                        "A mesclagem existe no Avaliado, "
                        "mas não no Base."
                    ),
                )

        tracker.update(
            10,
            f"Aba {sheet_index}/{len(common_sheets)} concluída.",
        )

    if progress_callback:
        progress_callback(
            100.0,
            "Comparação concluída.",
        )

    return differences


# ============================================================
# RELATÓRIO
# ============================================================

def generate_report(
    base_path,
    evaluated_path,
    differences,
    output_path,
    selected_options,
):
    wb = Workbook()

    ws_summary = wb.active
    ws_summary.title = "Resumo"

    ws_diff = wb.create_sheet(
        "Divergências"
    )

    title_font = Font(
        bold=True,
        size=14,
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )

    warning_fill = PatternFill(
        "solid",
        fgColor="FFF2CC",
    )

    ok_fill = PatternFill(
        "solid",
        fgColor="D9EAD3",
    )

    ws_summary["A1"] = (
        "Relatório de Comparação de Planilhas"
    )

    ws_summary["A1"].font = (
        title_font
    )

    ws_summary.merge_cells(
        "A1:D1"
    )

    summary_rows = [
        (
            "Arquivo Base",
            base_path,
        ),
        (
            "Arquivo Avaliado",
            evaluated_path,
        ),
        (
            "Data da comparação",
            datetime.now().strftime(
                "%d/%m/%Y %H:%M:%S"
            ),
        ),
        (
            "Total de divergências",
            len(differences),
        ),
    ]

    for idx, (label, value) in enumerate(
        summary_rows,
        start=3,
    ):
        ws_summary.cell(
            idx,
            1,
            label,
        ).font = Font(
            bold=True
        )

        ws_summary.cell(
            idx,
            2,
            value,
        )

    ws_summary["A9"] = "Resultado"
    ws_summary["A9"].font = Font(
        bold=True
    )

    if differences:
        ws_summary["B9"] = (
            f"Foram encontradas "
            f"{len(differences)} divergências."
        )

        ws_summary["B9"].fill = (
            warning_fill
        )

    else:
        ws_summary["B9"] = (
            "Nenhuma divergência encontrada "
            "para os parâmetros selecionados."
        )

        ws_summary["B9"].fill = (
            ok_fill
        )

    ws_summary["A11"] = (
        "Parâmetros avaliados"
    )

    ws_summary["A11"].font = Font(
        bold=True,
        size=11,
    )

    row = 12

    for label in selected_options:
        ws_summary.cell(
            row,
            1,
            "✓",
        )

        ws_summary.cell(
            row,
            2,
            label,
        )

        row += 1

    row += 2

    ws_summary.cell(
        row,
        1,
        "Categoria",
    )

    ws_summary.cell(
        row,
        2,
        "Quantidade",
    )

    for cell in (
        ws_summary.cell(row, 1),
        ws_summary.cell(row, 2),
    ):
        cell.font = header_font
        cell.fill = header_fill

    row += 1

    category_count = {}

    for item in differences:
        category = item["category"]

        category_count[category] = (
            category_count.get(
                category,
                0,
            )
            + 1
        )

    for category, count in sorted(
        category_count.items(),
        key=lambda x: (
            -x[1],
            x[0],
        ),
    ):
        ws_summary.cell(
            row,
            1,
            category,
        )

        ws_summary.cell(
            row,
            2,
            count,
        )

        row += 1

    headers = [
        "Planilha",
        "Local",
        "Categoria",
        "Base",
        "Avaliada",
        "Detalhes",
    ]

    for col, header in enumerate(
        headers,
        start=1,
    ):
        cell = ws_diff.cell(
            1,
            col,
            header,
        )

        cell.font = header_font
        cell.fill = header_fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for r, item in enumerate(
        differences,
        start=2,
    ):
        values = [
            item["sheet"],
            item["location"],
            item["category"],
            item["base"],
            item["evaluated"],
            item["details"],
        ]

        for c, value in enumerate(
            values,
            start=1,
        ):
            cell = ws_diff.cell(
                r,
                c,
                (
                    value
                    if value is not None
                    else ""
                ),
            )

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    ws_diff.freeze_panes = "A2"

    ws_diff.auto_filter.ref = (
        f"A1:F{max(ws_diff.max_row, 1)}"
    )

    widths = {
        "A": 24,
        "B": 32,
        "C": 34,
        "D": 55,
        "E": 55,
        "F": 75,
    }

    for col, width in widths.items():
        ws_diff.column_dimensions[
            col
        ].width = width

    ws_summary.column_dimensions[
        "A"
    ].width = 32

    ws_summary.column_dimensions[
        "B"
    ].width = 85

    wb.save(
        output_path
    )


# ============================================================
# INTERFACE
# ============================================================

def format_seconds(seconds):
    if seconds is None:
        return "--:--"

    seconds = max(
        0,
        int(seconds),
    )

    minutes, seconds = divmod(
        seconds,
        60,
    )

    hours, minutes = divmod(
        minutes,
        60,
    )

    if hours:
        return (
            f"{hours:02d}:"
            f"{minutes:02d}:"
            f"{seconds:02d}"
        )

    return (
        f"{minutes:02d}:"
        f"{seconds:02d}"
    )


class ComparatorFrame(ttk.Frame):
    """Comparador de referência adaptado para uma guia da aplicação principal."""

    def __init__(self, master, on_close=None):
        super().__init__(master)
        self.root = self.winfo_toplevel()
        self.on_close = on_close

        self.base_path = tk.StringVar()
        self.eval_path = tk.StringVar()

        self.status = tk.StringVar(
            value=(
                "Selecione os dois arquivos e "
                "os parâmetros que deseja avaliar."
            )
        )

        self.time_status = tk.StringVar(
            value=(
                "Tempo decorrido: 00:00 | "
                "Estimativa restante: --:--"
            )
        )

        self.option_vars = {}
        self.option_labels = {}

        self.progress_queue = queue.Queue()

        self.worker_thread = None
        self.start_time = None
        self.current_progress = 0.0

        self.build_ui()

    def build_ui(self):
        # ======================================================
        # ESTRUTURA PRINCIPAL
        # ======================================================
        root_container = ttk.Frame(
            self
        )

        root_container.pack(
            fill="both",
            expand=True,
        )

        # ------------------------------------------------------
        # ÁREA INFERIOR FIXA
        # Os botões e o progresso permanecem sempre visíveis.
        # ------------------------------------------------------
        bottom_frame = ttk.Frame(
            root_container,
            padding=(
                20,
                8,
                20,
                16,
            ),
        )

        bottom_frame.pack(
            side="bottom",
            fill="x",
        )

        action_frame = ttk.Frame(
            bottom_frame
        )

        action_frame.pack(
            fill="x",
            pady=(0, 8),
        )

        self.compare_button = ttk.Button(
            action_frame,
            text="Comparar e gerar relatório",
            command=self.run_comparison,
        )

        self.compare_button.pack(
            side="left"
        )

        ttk.Button(
            action_frame,
            text="Limpar arquivos",
            command=self.clear_files,
        ).pack(
            side="left",
            padx=(10, 0),
        )

        ttk.Button(
            action_frame,
            text="Fechar comparador",
            command=self.request_close,
        ).pack(
            side="right",
        )

        progress_frame = ttk.LabelFrame(
            bottom_frame,
            text="Execução",
            padding=10,
        )

        progress_frame.pack(
            fill="x",
        )

        self.progress_bar = ttk.Progressbar(
            progress_frame,
            orient="horizontal",
            mode="determinate",
            maximum=100,
        )

        self.progress_bar.pack(
            fill="x",
            pady=(0, 6),
        )

        ttk.Label(
            progress_frame,
            textvariable=self.status,
            wraplength=900,
        ).pack(
            anchor="w"
        )

        ttk.Label(
            progress_frame,
            textvariable=self.time_status,
        ).pack(
            anchor="w",
            pady=(3, 0),
        )

        # ------------------------------------------------------
        # ÁREA SUPERIOR ROLÁVEL
        # ------------------------------------------------------
        scroll_container = ttk.Frame(
            root_container
        )

        scroll_container.pack(
            side="top",
            fill="both",
            expand=True,
        )

        self.canvas = tk.Canvas(
            scroll_container,
            highlightthickness=0,
            borderwidth=0,
        )

        vertical_scrollbar = ttk.Scrollbar(
            scroll_container,
            orient="vertical",
            command=self.canvas.yview,
        )

        self.canvas.configure(
            yscrollcommand=vertical_scrollbar.set
        )

        vertical_scrollbar.pack(
            side="right",
            fill="y",
        )

        self.canvas.pack(
            side="left",
            fill="both",
            expand=True,
        )

        self.scrollable_frame = ttk.Frame(
            self.canvas,
            padding=(
                20,
                18,
                16,
                8,
            ),
        )

        self.canvas_window = self.canvas.create_window(
            (0, 0),
            window=self.scrollable_frame,
            anchor="nw",
        )

        self.scrollable_frame.bind(
            "<Configure>",
            self._update_scroll_region,
        )

        self.canvas.bind(
            "<Configure>",
            self._resize_canvas_content,
        )

        # Mouse wheel funciona somente quando o mouse está
        # sobre a área rolável.
        self.canvas.bind(
            "<Enter>",
            self._bind_mousewheel,
        )

        self.canvas.bind(
            "<Leave>",
            self._unbind_mousewheel,
        )

        main = self.scrollable_frame

        ttk.Label(
            main,
            text="Comparador de Planilhas Excel",
            font=(
                "Segoe UI",
                18,
                "bold",
            ),
        ).pack(
            anchor="w",
            pady=(0, 5),
        )

        ttk.Label(
            main,
            text=(
                "O sistema detecta deslocamentos causados por "
                "linhas/colunas inseridas ou removidas antes de comparar "
                "as células, reduzindo falsos erros em cascata."
            ),
            wraplength=880,
        ).pack(
            anchor="w",
            pady=(0, 18),
        )

        self.create_file_selector(
            main,
            "1. Planilha Base",
            self.base_path,
            self.select_base_file,
        )

        self.create_file_selector(
            main,
            "2. Planilha Avaliada",
            self.eval_path,
            self.select_eval_file,
        )

        options_frame = ttk.LabelFrame(
            main,
            text="3. Selecione o que deve ser avaliado",
            padding=12,
        )

        options_frame.pack(
            fill="both",
            expand=True,
            pady=(5, 12),
        )

        toolbar = ttk.Frame(
            options_frame
        )

        toolbar.pack(
            fill="x",
            pady=(0, 8),
        )

        ttk.Button(
            toolbar,
            text="Selecionar tudo",
            command=self.select_all,
        ).pack(
            side="left"
        )

        ttk.Button(
            toolbar,
            text="Desmarcar tudo",
            command=self.clear_all,
        ).pack(
            side="left",
            padx=(8, 0),
        )

        grid_container = ttk.Frame(
            options_frame
        )

        grid_container.pack(
            fill="both",
            expand=True,
        )

        # ======================================================
        # CONTEÚDO
        # ======================================================
        content_frame = ttk.LabelFrame(
            grid_container,
            text="Conteúdo",
            padding=8,
        )

        content_frame.grid(
            row=0,
            column=0,
            sticky="nsew",
            padx=(0, 6),
            pady=6,
        )

        self.add_option(
            content_frame,
            "value",
            "Valor alterado",
        )

        self.add_option(
            content_frame,
            "formula",
            "Fórmula alterada",
        )

        self.add_option(
            content_frame,
            "removed_content",
            "Conteúdo removido",
        )

        self.add_option(
            content_frame,
            "added_content",
            "Conteúdo adicionado",
        )

        # ======================================================
        # FORMATAÇÃO
        # ======================================================
        formatting_frame = ttk.LabelFrame(
            grid_container,
            text="Formatação",
            padding=8,
        )

        formatting_frame.grid(
            row=0,
            column=1,
            sticky="nsew",
            padx=(6, 0),
            pady=6,
        )

        self.add_option(
            formatting_frame,
            "font",
            "Fonte",
        )

        self.add_option(
            formatting_frame,
            "fill",
            "Preenchimento",
        )

        self.add_option(
            formatting_frame,
            "border",
            "Borda",
        )

        self.add_option(
            formatting_frame,
            "alignment",
            "Alinhamento",
        )

        self.add_option(
            formatting_frame,
            "number_format",
            "Formato numérico",
        )

        self.add_option(
            formatting_frame,
            "protection",
            "Proteção da célula",
        )

        # ======================================================
        # ESTRUTURA
        # ======================================================
        structure_frame = ttk.LabelFrame(
            grid_container,
            text="Estrutura",
            padding=8,
        )

        structure_frame.grid(
            row=1,
            column=0,
            sticky="nsew",
            padx=(0, 6),
            pady=6,
        )

        self.add_option(
            structure_frame,
            "added_rows",
            "Linhas adicionadas",
        )

        self.add_option(
            structure_frame,
            "removed_rows",
            "Linhas removidas",
        )

        self.add_option(
            structure_frame,
            "added_columns",
            "Colunas adicionadas",
        )

        self.add_option(
            structure_frame,
            "removed_columns",
            "Colunas removidas",
        )

        self.add_option(
            structure_frame,
            "row_height",
            "Altura das linhas",
        )

        self.add_option(
            structure_frame,
            "column_width",
            "Largura das colunas",
        )

        self.add_option(
            structure_frame,
            "row_hidden",
            "Linhas ocultas",
        )

        self.add_option(
            structure_frame,
            "column_hidden",
            "Colunas ocultas",
        )

        self.add_option(
            structure_frame,
            "merged_cells",
            "Células mescladas",
        )

        # ======================================================
        # PLANILHA / ABAS
        # ======================================================
        sheet_frame = ttk.LabelFrame(
            grid_container,
            text="Planilha / Abas",
            padding=8,
        )

        sheet_frame.grid(
            row=1,
            column=1,
            sticky="nsew",
            padx=(6, 0),
            pady=6,
        )

        self.add_option(
            sheet_frame,
            "sheets",
            "Abas adicionadas ou removidas",
        )

        self.add_option(
            sheet_frame,
            "sheet_state",
            "Estado da aba (visível/oculta)",
        )

        self.add_option(
            sheet_frame,
            "sheet_protection",
            "Proteção da planilha",
        )

        self.add_option(
            sheet_frame,
            "freeze_panes",
            "Painel congelado",
        )

        grid_container.columnconfigure(
            0,
            weight=1,
            uniform="options",
        )

        grid_container.columnconfigure(
            1,
            weight=1,
            uniform="options",
        )

        grid_container.rowconfigure(
            0,
            weight=1,
        )

        grid_container.rowconfigure(
            1,
            weight=1,
        )

    def request_close(self):
        if self.worker_thread and self.worker_thread.is_alive():
            messagebox.showwarning(
                APP_TITLE,
                "Aguarde a comparação em andamento antes de fechar.",
            )
            return
        self._unbind_mousewheel()
        if self.on_close:
            self.on_close()

    # ==========================================================
    # ROLAGEM / RESPONSIVIDADE
    # ==========================================================

    def _update_scroll_region(self, event=None):
        self.canvas.configure(
            scrollregion=self.canvas.bbox("all")
        )

    def _resize_canvas_content(self, event):
        # Faz o frame interno acompanhar a largura disponível.
        self.canvas.itemconfigure(
            self.canvas_window,
            width=event.width,
        )

    def _bind_mousewheel(self, event=None):
        self.root.bind_all(
            "<MouseWheel>",
            self._on_mousewheel,
        )

    def _unbind_mousewheel(self, event=None):
        self.root.unbind_all(
            "<MouseWheel>"
        )

    def _on_mousewheel(self, event):
        if event.delta:
            self.canvas.yview_scroll(
                int(-1 * (event.delta / 120)),
                "units",
            )

    # ==========================================================
    # OPÇÕES
    # ==========================================================

    def add_option(
        self,
        parent,
        key,
        label,
        default=True,
    ):
        var = tk.BooleanVar(
            value=default
        )

        self.option_vars[key] = var
        self.option_labels[key] = label

        ttk.Checkbutton(
            parent,
            text=label,
            variable=var,
        ).pack(
            anchor="w",
            pady=2,
        )

    def create_file_selector(
        self,
        parent,
        label_text,
        variable,
        command,
    ):
        frame = ttk.Frame(
            parent
        )

        frame.pack(
            fill="x",
            pady=(0, 12),
        )

        ttk.Label(
            frame,
            text=label_text,
            font=(
                "Segoe UI",
                10,
                "bold",
            ),
        ).pack(
            anchor="w"
        )

        row = ttk.Frame(
            frame
        )

        row.pack(
            fill="x",
            pady=(5, 0),
        )

        ttk.Entry(
            row,
            textvariable=variable,
        ).pack(
            side="left",
            fill="x",
            expand=True,
        )

        ttk.Button(
            row,
            text="Selecionar...",
            command=command,
        ).pack(
            side="left",
            padx=(8, 0),
        )

    # ==========================================================
    # ARQUIVOS
    # ==========================================================

    def select_base_file(self):
        path = filedialog.askopenfilename(
            title="Selecionar planilha Base",
            filetypes=SUPPORTED_EXTENSIONS,
        )

        if path:
            self.base_path.set(
                path
            )

    def select_eval_file(self):
        path = filedialog.askopenfilename(
            title="Selecionar planilha Avaliada",
            filetypes=SUPPORTED_EXTENSIONS,
        )

        if path:
            self.eval_path.set(
                path
            )

    def select_all(self):
        for var in self.option_vars.values():
            var.set(
                True
            )

    def clear_all(self):
        for var in self.option_vars.values():
            var.set(
                False
            )

    def clear_files(self):
        self.base_path.set(
            ""
        )

        self.eval_path.set(
            ""
        )

        self.progress_bar[
            "value"
        ] = 0

        self.status.set(
            "Selecione os dois arquivos e "
            "os parâmetros que deseja avaliar."
        )

        self.time_status.set(
            "Tempo decorrido: 00:00 | "
            "Estimativa restante: --:--"
        )

        # Retorna a área rolável para o topo.
        self.canvas.yview_moveto(
            0
        )

    def get_options(self):
        return {
            key: var.get()
            for key, var
            in self.option_vars.items()
        }

    def get_selected_option_labels(self):
        return [
            self.option_labels[key]
            for key, var
            in self.option_vars.items()
            if var.get()
        ]

    # ==========================================================
    # EXECUÇÃO
    # ==========================================================

    def run_comparison(self):
        base = (
            self.base_path.get().strip()
        )

        evaluated = (
            self.eval_path.get().strip()
        )

        if not base or not evaluated:
            messagebox.showwarning(
                APP_TITLE,
                (
                    "Selecione a planilha Base e "
                    "a planilha Avaliada."
                ),
            )

            return

        if (
            not os.path.isfile(base)
            or not os.path.isfile(evaluated)
        ):
            messagebox.showerror(
                APP_TITLE,
                (
                    "Um dos arquivos selecionados "
                    "não foi encontrado."
                ),
            )

            return

        if (
            os.path.abspath(base)
            == os.path.abspath(evaluated)
        ):
            messagebox.showwarning(
                APP_TITLE,
                (
                    "Selecione dois arquivos "
                    "diferentes."
                ),
            )

            return

        options = (
            self.get_options()
        )

        if not any(
            options.values()
        ):
            messagebox.showwarning(
                APP_TITLE,
                (
                    "Selecione pelo menos um "
                    "parâmetro para comparação."
                ),
            )

            return

        self.compare_button.config(
            state="disabled"
        )

        self.progress_bar[
            "value"
        ] = 0

        self.current_progress = 0.0
        self.start_time = time.time()

        self.status.set(
            "Iniciando comparação..."
        )

        self.time_status.set(
            "Tempo decorrido: 00:00 | "
            "Estimativa restante: calculando..."
        )

        selected_labels = (
            self.get_selected_option_labels()
        )

        self.worker_thread = threading.Thread(
            target=self.comparison_worker,
            args=(
                base,
                evaluated,
                options,
                selected_labels,
            ),
            daemon=True,
        )

        self.worker_thread.start()

        self.poll_worker_queue()
        self.update_timer()

    def comparison_worker(
        self,
        base,
        evaluated,
        options,
        selected_labels,
    ):
        try:
            def progress_callback(
                percent,
                message,
            ):
                self.progress_queue.put(
                    (
                        "progress",
                        percent,
                        message,
                    )
                )

            differences = compare_workbooks(
                base,
                evaluated,
                options,
                progress_callback=progress_callback,
            )

            self.progress_queue.put(
                (
                    "done",
                    base,
                    evaluated,
                    differences,
                    selected_labels,
                )
            )

        except Exception as exc:
            self.progress_queue.put(
                (
                    "error",
                    exc,
                )
            )

    def poll_worker_queue(self):
        try:
            while True:
                item = (
                    self.progress_queue.get_nowait()
                )

                event_type = item[0]

                if event_type == "progress":
                    _, percent, message = item

                    self.current_progress = max(
                        self.current_progress,
                        percent,
                    )

                    self.progress_bar[
                        "value"
                    ] = self.current_progress

                    self.status.set(
                        message
                    )

                elif event_type == "done":
                    (
                        _,
                        base,
                        evaluated,
                        differences,
                        selected_labels,
                    ) = item

                    self.progress_bar[
                        "value"
                    ] = 100

                    self.current_progress = 100

                    self.finish_comparison(
                        base,
                        evaluated,
                        differences,
                        selected_labels,
                    )

                    return

                elif event_type == "error":
                    _, exc = item

                    self.compare_button.config(
                        state="normal"
                    )

                    self.status.set(
                        "A comparação foi interrompida por um erro."
                    )

                    messagebox.showerror(
                        APP_TITLE,
                        (
                            "Ocorreu um erro durante "
                            f"a comparação:\n\n{exc}"
                        ),
                    )

                    return

        except queue.Empty:
            pass

        if (
            self.worker_thread
            and self.worker_thread.is_alive()
        ):
            self.root.after(
                150,
                self.poll_worker_queue,
            )

    def update_timer(self):
        if self.start_time is None:
            return

        elapsed = (
            time.time()
            - self.start_time
        )

        remaining = None

        if (
            self.current_progress
            > 1.0
            and self.current_progress
            < 100.0
        ):
            total_estimated = (
                elapsed
                / (
                    self.current_progress
                    / 100.0
                )
            )

            remaining = (
                total_estimated
                - elapsed
            )

        self.time_status.set(
            (
                f"Tempo decorrido: "
                f"{format_seconds(elapsed)} | "
                f"Estimativa restante: "
                f"{format_seconds(remaining)}"
            )
        )

        if (
            self.worker_thread
            and self.worker_thread.is_alive()
        ):
            self.root.after(
                1000,
                self.update_timer,
            )

    def finish_comparison(
        self,
        base,
        evaluated,
        differences,
        selected_labels,
    ):
        elapsed = (
            time.time()
            - self.start_time
            if self.start_time
            else 0
        )

        self.status.set(
            (
                "Comparação concluída. "
                f"{len(differences)} divergências encontradas."
            )
        )

        self.time_status.set(
            (
                f"Tempo total: "
                f"{format_seconds(elapsed)}"
            )
        )

        default_name = (
            "Relatorio_Comparacao_"
            + datetime.now().strftime(
                "%Y%m%d_%H%M%S"
            )
            + ".xlsx"
        )

        output_path = filedialog.asksaveasfilename(
            title="Salvar relatório",
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[
                (
                    "Planilha Excel",
                    "*.xlsx",
                )
            ],
        )

        if not output_path:
            self.compare_button.config(
                state="normal"
            )

            return

        try:
            self.status.set(
                "Gerando relatório..."
            )

            self.root.update_idletasks()

            generate_report(
                base,
                evaluated,
                differences,
                output_path,
                selected_labels,
            )

            self.status.set(
                (
                    f"Concluído: "
                    f"{len(differences)} divergências. "
                    f"Relatório salvo em: {output_path}"
                )
            )

            messagebox.showinfo(
                APP_TITLE,
                (
                    "Comparação concluída.\n\n"
                    f"Tempo total: "
                    f"{format_seconds(elapsed)}\n"
                    f"Parâmetros analisados: "
                    f"{len(selected_labels)}\n"
                    f"Divergências encontradas: "
                    f"{len(differences)}\n\n"
                    f"Relatório:\n{output_path}"
                ),
            )

            try:
                os.startfile(
                    output_path
                )
            except Exception:
                pass

        except PermissionError:
            messagebox.showerror(
                APP_TITLE,
                (
                    "Não foi possível salvar o relatório.\n\n"
                    "Feche o arquivo no Excel e tente novamente."
                ),
            )

        except Exception as exc:
            messagebox.showerror(
                APP_TITLE,
                (
                    "Erro ao gerar o relatório:\n\n"
                    f"{exc}"
                ),
            )

        finally:
            self.compare_button.config(
                state="normal"
            )
