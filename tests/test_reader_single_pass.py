from datetime import date, datetime
from pathlib import Path
import xml.etree.ElementTree as ET
import zipfile

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from app.excel import reader


def _assert_readers_are_typed_equal(path: Path) -> None:
    previous = reader._read_fast_repeated_find(path)
    optimized = reader._read_fast(path)
    assert reader.read_workbook(path) == optimized
    assert previous.keys() == optimized.keys()
    for title in previous:
        assert previous[title].keys() == optimized[title].keys()
        for coordinate, value in previous[title].items():
            assert type(value) is type(optimized[title][coordinate])
            assert value == optimized[title][coordinate]


@pytest.mark.parametrize("use_1904_epoch", [False, True])
def test_single_pass_matches_previous_reader_for_cell_types_and_sheets(
    tmp_path: Path, use_1904_epoch: bool
) -> None:
    path = tmp_path / f"types-{use_1904_epoch}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Dados ç漢字"
    values = [
        12, 2.5, 0, False, True, "texto", "Unicode ç漢字", "", None,
        date(2024, 2, 29), datetime(2024, 2, 29, 12, 34, 56), "=SUM(A1:A3)",
    ]
    for row, value in enumerate(values, 1):
        sheet.cell(row=row, column=1, value=value)
    workbook.create_sheet("Aba vazia")
    workbook.create_sheet("Terceira")
    if use_1904_epoch:
        workbook.epoch = CALENDAR_MAC_1904
    workbook.save(path)
    workbook.close()

    _assert_readers_are_typed_equal(path)


def _rewrite_xlsx(path: Path, replacements: dict[str, bytes]) -> None:
    rewritten = path.with_suffix(".rewritten.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w") as target:
        for item in source.infolist():
            if item.filename not in replacements:
                target.writestr(item, source.read(item.filename))
        for name, content in replacements.items():
            target.writestr(name, content)
    rewritten.replace(path)


def test_single_pass_matches_previous_for_shared_inline_rich_and_shared_formula(
    tmp_path: Path,
) -> None:
    path = tmp_path / "xml-features.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "placeholder"
    sheet["A2"] = "placeholder"
    sheet["B1"] = "placeholder"
    sheet["C1"] = "=A1+1"
    sheet["C2"] = 0
    sheet["D1"] = "#DIV/0!"
    sheet["D1"].data_type = "e"
    workbook.save(path)
    workbook.close()

    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    cells = {cell.get("r"): cell for cell in root.iter(reader._tag("c"))}

    cells["A1"].set("t", "s")
    cells["A1"].clear()
    cells["A1"].set("r", "A1")
    cells["A1"].set("t", "s")
    ET.SubElement(cells["A1"], reader._tag("v")).text = "0"

    cells["A2"].clear()
    cells["A2"].set("r", "A2")
    cells["A2"].set("t", "inlineStr")
    inline = ET.SubElement(cells["A2"], reader._tag("is"))
    for text in ("inline ", "rico"):
        run = ET.SubElement(inline, reader._tag("r"))
        ET.SubElement(run, reader._tag("t")).text = text

    formula = cells["C1"].find("m:f", reader._NS)
    assert formula is not None
    formula.set("t", "shared")
    formula.set("si", "7")
    formula.set("ref", "C1:C2")
    cells["C2"].clear()
    cells["C2"].set("r", "C2")
    follower = ET.SubElement(cells["C2"], reader._tag("f"))
    follower.set("t", "shared")
    follower.set("si", "7")

    shared_strings = (
        f'<sst xmlns="{reader._MAIN_NS}" count="1" uniqueCount="1">'
        "<si><r><t>shared </t></r><r><t>rico ç漢字</t></r></si></sst>"
    ).encode()
    _rewrite_xlsx(
        path,
        {
            "xl/worksheets/sheet1.xml": ET.tostring(root),
            "xl/sharedStrings.xml": shared_strings,
        },
    )

    _assert_readers_are_typed_equal(path)
    snapshot = reader.read_workbook(path)["Sheet"]
    assert snapshot["A1"] == "shared rico ç漢字"
    assert snapshot["A2"] == "inline rico"
    assert snapshot["C1"] == "=A1+1"
    assert snapshot["C2"] == "=A2+1"
    assert snapshot["D1"] == "#DIV/0!"


@pytest.mark.parametrize("formula_type", ["array", "dataTable"])
def test_array_and_data_table_keep_openpyxl_fallback(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path, formula_type: str
) -> None:
    cell = ET.fromstring(
        f'<c xmlns="{reader._MAIN_NS}" r="A1"><f t="{formula_type}">A2</f></c>'
    )
    arguments = (cell, "A1", [], set(), reader.CALENDAR_WINDOWS_1900, {})
    for cell_reader in (reader._cell_value_repeated_find, reader._cell_value):
        with pytest.raises(reader._FastReaderUnsupported, match=formula_type):
            cell_reader(*arguments)

    path = tmp_path / f"{formula_type}.xlsx"
    path.touch()
    expected = {"Fallback": {"A1": f"{formula_type} preserved"}}
    monkeypatch.setattr(
        reader,
        "_read_fast",
        lambda unused: (_ for _ in ()).throw(
            reader._FastReaderUnsupported(f"fórmula {formula_type} em A1")
        ),
    )
    monkeypatch.setattr(reader, "_read_openpyxl", lambda unused: expected)

    assert reader.read_workbook(path) is expected


def test_single_pass_preserves_first_duplicate_child_semantics() -> None:
    cell = ET.fromstring(
        f'<c xmlns="{reader._MAIN_NS}" r="A1"><v>1</v><v>2</v></c>'
    )
    arguments = (cell, "A1", [], set(), reader.CALENDAR_WINDOWS_1900, {})

    assert reader._cell_value(*arguments) == reader._cell_value_repeated_find(*arguments) == 1
