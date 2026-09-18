from pathlib import Path

from openpyxl import load_workbook
import pytest

from app.database import Database
from app.report_service import ReportService


@pytest.fixture
def populated_database(tmp_path: Path):
    with Database(tmp_path / "audit.db") as database:
        database.initialize()
        connection = database.connection
        spreadsheet_id = connection.execute(
            """INSERT INTO planilha
               (drive_item_id, nome_atual, site_id, drive_id, caminho_sharepoint)
               VALUES ('item-1', 'CQL028.xlsx', 'site-1', 'drive-1', '/docs/CQL028.xlsx')"""
        ).lastrowid
        execution_id = connection.execute(
            """INSERT INTO execucao_auditoria
               (codigo_execucao, planilha_id, status, fim, versoes_processadas,
                alteracoes_encontradas)
               VALUES ('AUD-1', ?, 'CONCLUIDA', CURRENT_TIMESTAMP, 1, 2)""",
            (spreadsheet_id,),
        ).lastrowid
        version_id = connection.execute(
            """INSERT INTO versao_processada
               (planilha_id, versao_anterior_id, versao_anterior_numero,
                versao_atual_id, versao_atual_numero, data_hora_versao, autor,
                comentario, quantidade_alteracoes, status, execucao_id)
               VALUES (?, 'v1', '0.1', 'v2', '0.2', '2026-09-16T10:00:00Z',
                       'Maria', 'Ajuste', 2, 'PROCESSADA', ?)""",
            (spreadsheet_id, execution_id),
        ).lastrowid
        connection.executemany(
            """INSERT INTO alteracao
               (versao_processada_id, planilha_id, tipo, aba, endereco,
                valor_anterior, valor_novo) VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (
                (version_id, spreadsheet_id, "ADD", "Dados", "A1", None, "novo"),
                (version_id, spreadsheet_id, "MOD", "Dados", "B2", "1", "2"),
            ),
        )
        connection.commit()
        yield database, spreadsheet_id


def test_report_contains_official_sheets_filters_and_database_data(
    populated_database, tmp_path: Path
) -> None:
    database, spreadsheet_id = populated_database

    report = ReportService(database.connection, tmp_path / "reports").generate(
        spreadsheet_id
    )
    workbook = load_workbook(report)

    assert report.name.startswith("CQL028__")
    assert report.name.endswith("_Trilha_Auditoria.xlsx")
    assert workbook.sheetnames == ["RESUMO", "VERSOES", "TRILHA"]
    assert workbook["RESUMO"]["B2"].value == "CQL028.xlsx"
    assert workbook["RESUMO"]["B9"].value == 2
    assert workbook["VERSOES"]["A2"].value == "0.1"
    assert workbook["TRILHA"]["I2"].value == "ADD"
    assert all(sheet.auto_filter.ref for sheet in workbook.worksheets)


def test_report_can_be_regenerated_and_rejects_unknown_spreadsheet(
    populated_database, tmp_path: Path
) -> None:
    database, spreadsheet_id = populated_database
    service = ReportService(database.connection, tmp_path)

    first = service.generate(spreadsheet_id)
    first.unlink()
    second = service.generate(spreadsheet_id)

    assert second.is_file()
    with pytest.raises(ValueError, match="não encontrada"):
        service.generate(999)


def test_report_is_deterministically_regenerated_from_database_only(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "audit.db"
    historical_directory = tmp_path / "historical_versions"
    historical_directory.mkdir()
    (historical_directory / "0.1.xlsx").write_bytes(b"temporary historical file")
    (historical_directory / "0.2.xlsx").write_bytes(b"temporary historical file")

    with Database(database_path) as database:
        database.initialize()
        connection = database.connection
        spreadsheet_id = connection.execute(
            """INSERT INTO planilha
               (drive_item_id, nome_atual, site_id, drive_id, caminho_sharepoint)
               VALUES ('stable-item-id', 'Auditoria Oficial.xlsx', 'site-oficial',
                       'drive-oficial', '/documentos/Auditoria Oficial.xlsx')"""
        ).lastrowid
        execution_id = connection.execute(
            """INSERT INTO execucao_auditoria
               (codigo_execucao, planilha_id, inicio, fim, status,
                checkpoint_inicial, versao_final, versoes_processadas,
                alteracoes_encontradas)
               VALUES ('AUD-RELATORIO', ?, '2026-09-16T09:00:00Z',
                       '2026-09-16T09:05:00Z', 'CONCLUIDA', '0.1', '0.4', 3, 3)""",
            (spreadsheet_id,),
        ).lastrowid
        version_rows = (
            (
                spreadsheet_id,
                "version-1",
                "0.1",
                "version-2",
                "0.2",
                "2026-09-16T08:01:00Z",
                "Ana",
                "Inclusão",
                1,
                "PROCESSADA",
                execution_id,
            ),
            (
                spreadsheet_id,
                "version-2",
                "0.2",
                "version-3",
                "0.3",
                "2026-09-16T08:02:00Z",
                "Bruno",
                "Sem diferenças",
                0,
                "SEM_ALTERACOES",
                execution_id,
            ),
            (
                spreadsheet_id,
                "version-3",
                "0.3",
                "version-4",
                "0.4",
                "2026-09-16T08:03:00Z",
                "Carla",
                "Fórmula e remoção",
                2,
                "PROCESSADA",
                execution_id,
            ),
        )
        version_ids = []
        for row in version_rows:
            version_ids.append(
                connection.execute(
                    """INSERT INTO versao_processada
                       (planilha_id, versao_anterior_id, versao_anterior_numero,
                        versao_atual_id, versao_atual_numero, data_hora_versao,
                        autor, comentario, quantidade_alteracoes, status, execucao_id)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    row,
                ).lastrowid
            )
        connection.executemany(
            """INSERT INTO alteracao
               (versao_processada_id, planilha_id, tipo, aba, endereco,
                valor_anterior, valor_novo) VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (
                (version_ids[0], spreadsheet_id, "ADD", "Dados", "A1", None, "novo"),
                (
                    version_ids[2],
                    spreadsheet_id,
                    "MOD",
                    "Cálculos",
                    "C3",
                    "=SUM(A1:A2)",
                    "=SUM(A1:A3)",
                ),
                (version_ids[2], spreadsheet_id, "DEL", "Dados", "B2", "removido", None),
            ),
        )
        connection.commit()

        for historical_file in historical_directory.iterdir():
            historical_file.unlink()
        historical_directory.rmdir()

        first_report = ReportService(connection, tmp_path / "first").generate(
            spreadsheet_id
        )
        second_report = ReportService(connection, tmp_path / "second").generate(
            spreadsheet_id
        )

        first = load_workbook(first_report, data_only=False)
        second = load_workbook(second_report, data_only=False)
        try:
            assert first.sheetnames == second.sheetnames == [
                "RESUMO",
                "VERSOES",
                "TRILHA",
            ]
            for sheet_name in first.sheetnames:
                first_rows = list(first[sheet_name].iter_rows(values_only=True))
                second_rows = list(second[sheet_name].iter_rows(values_only=True))
                assert first_rows == second_rows

            summary = dict(first["RESUMO"].iter_rows(min_row=2, values_only=True))
            assert summary == {
                "Planilha": "Auditoria Oficial.xlsx",
                "Identidade técnica": "site-oficial/drive-oficial/stable-item-id",
                "DriveItem ID": "stable-item-id",
                "Primeira versão": "0.1",
                "Última versão": "0.4",
                "Última execução": "2026-09-16T09:00:00Z",
                "Versões processadas": 3,
                "Total de alterações": 3,
                "ADD": 1,
                "MOD": 1,
                "DEL": 1,
            }
            assert list(first["VERSOES"].iter_rows(min_row=2, values_only=True)) == [
                (
                    "0.1",
                    "0.2",
                    "2026-09-16T08:01:00Z",
                    "Ana",
                    "Inclusão",
                    "PROCESSADA",
                    1,
                ),
                (
                    "0.2",
                    "0.3",
                    "2026-09-16T08:02:00Z",
                    "Bruno",
                    "Sem diferenças",
                    "SEM_ALTERACOES",
                    0,
                ),
                (
                    "0.3",
                    "0.4",
                    "2026-09-16T08:03:00Z",
                    "Carla",
                    "Fórmula e remoção",
                    "PROCESSADA",
                    2,
                ),
            ]
            trail_rows = list(first["TRILHA"].iter_rows(min_row=2, values_only=True))
            assert [row[8] for row in trail_rows] == ["ADD", "MOD", "DEL"]
            assert trail_rows[1][6:] == (
                "Cálculos",
                "C3",
                "MOD",
                "=SUM(A1:A2)",
                "=SUM(A1:A3)",
            )
            assert trail_rows[2][6:] == (
                "Dados",
                "B2",
                "DEL",
                "removido",
                None,
            )
            assert all(sheet.auto_filter.ref for sheet in first.worksheets)
            assert all(sheet.freeze_panes == "A2" for sheet in first.worksheets)
        finally:
            first.close()
            second.close()

        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
