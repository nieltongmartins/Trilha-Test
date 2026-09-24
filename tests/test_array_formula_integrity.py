from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from app.audit_service import AuditService
from app.database import Database
from app.excel.comparator import compare_snapshots
from app.excel.formula_values import normalize_formula_value
from app.excel.reader import ConsecutiveWorkbookReader, read_workbook
from app.models import ChangeType
from app.report_service import ReportService
from app.sources import LocalSource, SpreadsheetInfo, VersionInfo


def _array(ref: str = "D2:D10", text: str = "=ROW(D2:D10)") -> ArrayFormula:
    return ArrayFormula(ref=ref, text=text)


def _write_array_workbook(path: Path, *, ref: str, text: str) -> None:
    workbook = Workbook()
    workbook.active["D2"] = ArrayFormula(ref=ref, text=text)
    workbook.save(path)
    workbook.close()


def test_common_formulas_compare_without_changing_existing_representation() -> None:
    assert normalize_formula_value("=SUM(A1:A2)") == "=SUM(A1:A2)"
    assert compare_snapshots(
        {"Dados": {"A1": "=SUM(A1:A2)"}},
        {"Dados": {"A1": "=SUM(A1:A2)"}},
    ) == []

    changes = compare_snapshots(
        {"Dados": {"A1": "=SUM(A1:A2)"}},
        {"Dados": {"A1": "=SUM(A1:A3)"}},
    )
    assert len(changes) == 1
    assert changes[0].change_type is ChangeType.MOD


def test_array_formula_comparison_uses_semantics_not_object_identity() -> None:
    first = _array()
    second = _array()
    assert first is not second
    assert compare_snapshots({"Dados": {"D2": first}}, {"Dados": {"D2": second}}) == []

    formula_change = compare_snapshots(
        {"Dados": {"D2": _array()}},
        {"Dados": {"D2": _array(text="=ROW(D2:D10)+1")}},
    )
    range_change = compare_snapshots(
        {"Dados": {"D2": _array()}},
        {"Dados": {"D2": _array(ref="D2:D11")}},
    )
    assert len(formula_change) == len(range_change) == 1
    assert formula_change[0].previous_value == (
        'ARRAYFORMULA|{"formula":"=ROW(D2:D10)","ref":"D2:D10"}'
    )
    assert "object at 0x" not in formula_change[0].previous_value


def test_other_special_formula_and_non_formula_values_are_stable() -> None:
    data_table = DataTableFormula("A1:B3", dt2D=True, r1="A1", r2="B1")
    normalized = normalize_formula_value(data_table)
    assert normalized.startswith("DATATABLEFORMULA|")
    assert '"ref":"A1:B3"' in normalized
    assert '"dt2D":true' in normalized

    for value in (None, 0, False, ""):
        assert normalize_formula_value(value) is value
    assert compare_snapshots(
        {"Dados": {"A1": None, "A2": 0, "A3": False}},
        {"Dados": {"A1": "", "A2": False, "A3": 0}},
    ) == compare_snapshots(
        {"Dados": {"A2": 0, "A3": False}},
        {"Dados": {"A1": "", "A2": False, "A3": 0}},
    )


def test_real_array_formula_file_is_normalized_by_openpyxl_fallback(tmp_path: Path) -> None:
    path = tmp_path / "array.xlsx"
    _write_array_workbook(path, ref="D2:D10", text="=ROW(D2:D10)")

    raw = load_workbook(path, read_only=True, data_only=False)
    try:
        assert isinstance(raw.active["D2"].value, ArrayFormula)
        assert raw.active["D2"].value.text == "=ROW(D2:D10)"
        assert raw.active["D2"].value.ref == "D2:D10"
    finally:
        raw.close()

    reader = ConsecutiveWorkbookReader()
    snapshot = reader.read(path)
    assert reader.last_metrics.fallback_used is True
    assert snapshot["Sheet"]["D2"] == (
        'ARRAYFORMULA|{"formula":"=ROW(D2:D10)","ref":"D2:D10"}'
    )
    assert snapshot == read_workbook(path)


def test_array_formula_audit_sqlite_and_report_never_use_object_repr(tmp_path: Path) -> None:
    previous_path = tmp_path / "1.xlsx"
    current_path = tmp_path / "2.xlsx"
    _write_array_workbook(previous_path, ref="D2:D10", text="=ROW(D2:D10)")
    _write_array_workbook(current_path, ref="D2:D10", text="=ROW(D2:D10)+1")
    spreadsheet = SpreadsheetInfo("site", "drive", "item", "array.xlsx")
    versions = [
        (VersionInfo("v1", "1"), previous_path),
        (VersionInfo("v2", "2"), current_path),
    ]
    source = LocalSource(
        [spreadsheet],
        {(spreadsheet.site_id, spreadsheet.drive_id, spreadsheet.drive_item_id): versions},
    )

    with Database(tmp_path / "audit.db") as database:
        database.initialize()
        result = AuditService(database, source).audit(spreadsheet)
        assert result.changes == 1
        row = database.connection.execute(
            "SELECT valor_anterior, valor_novo FROM alteracao"
        ).fetchone()
        assert row["valor_anterior"].startswith("ARRAYFORMULA|")
        assert row["valor_novo"].startswith("ARRAYFORMULA|")
        assert "object at 0x" not in row["valor_anterior"] + row["valor_novo"]

        spreadsheet_id = database.connection.execute("SELECT id FROM planilha").fetchone()[0]
        report = ReportService(database.connection, tmp_path / "reports").generate(
            spreadsheet_id
        )

    workbook = load_workbook(report, read_only=True, data_only=False)
    try:
        old_value = workbook["Alterações_001"]["J2"].value
        new_value = workbook["Alterações_001"]["K2"].value
        assert old_value.startswith("ARRAYFORMULA|")
        assert "=ROW(D2:D10)" in old_value
        assert "=ROW(D2:D10)+1" in new_value
        assert "ArrayFormula object at" not in old_value + new_value
    finally:
        workbook.close()
