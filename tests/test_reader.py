from pathlib import Path
import shutil
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from app.excel.reader import ConsecutiveWorkbookReader, read_workbook
from app.excel.comparator import compare_snapshots
from app.excel import reader


def _rewrite_member(path: Path, name: str, transform) -> None:
    rewritten = path.with_suffix(".rewritten.xlsx")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(rewritten, "w") as target:
        for item in source.infolist():
            payload = source.read(item.filename)
            target.writestr(item, transform(payload) if item.filename == name else payload)
    rewritten.replace(path)


def test_reader_preserves_formulas_values_and_sheets(
    cql028_versions: Path,
) -> None:
    snapshot = read_workbook(cql028_versions / "0.85.xlsx")

    assert list(snapshot) == ["Resumo", "Detalhes"]
    assert snapshot["Resumo"] == {
        "A1": 15,
        "B2": "OK",
        "C3": "Pendente",
        "D4": 0,
        "E5": False,
        "F6": "=SUM(A1:A20)",
    }
    assert snapshot["Detalhes"] == {"B2": "Nova informação"}


def test_reader_ignores_explicitly_empty_cells(cql028_versions: Path) -> None:
    snapshot = read_workbook(cql028_versions / "0.84.xlsx")

    # B2 é gravada como None pela fixture programática; openpyxl a materializa
    # como uma célula vazia e o snapshot não deve preservá-la.
    assert "B2" not in snapshot["Resumo"]
    assert "Detalhes" in snapshot
    assert snapshot["Detalhes"] == {}


def test_reader_rejects_non_xlsx_file(tmp_path: Path) -> None:
    path = tmp_path / "arquivo.xls"
    path.touch()

    with pytest.raises(ValueError, match="somente arquivos .xlsx"):
        read_workbook(path)


def test_incremental_reader_reuses_only_cryptographically_equal_sheets(
    cql028_versions: Path,
) -> None:
    incremental = ConsecutiveWorkbookReader()
    previous = incremental.read(cql028_versions / "0.86.xlsx")
    assert incremental.last_metrics.worksheets_reused == 0

    current = incremental.read(cql028_versions / "0.87.xlsx")

    assert current == read_workbook(cql028_versions / "0.87.xlsx")
    assert incremental.last_metrics.worksheets_reused == 3
    assert incremental.last_metrics.cells_parsed == 0
    assert all(previous[name] is current[name] for name in previous)


def test_incremental_reader_invalidates_changed_worksheet(
    cql028_versions: Path,
) -> None:
    incremental = ConsecutiveWorkbookReader()
    previous = incremental.read(cql028_versions / "0.84.xlsx")
    current = incremental.read(cql028_versions / "0.85.xlsx")

    assert current == read_workbook(cql028_versions / "0.85.xlsx")
    assert incremental.last_metrics.worksheets_reused == 0
    assert previous["Resumo"] is not current["Resumo"]


def test_incremental_reader_reuses_cryptographically_equal_rows(tmp_path: Path) -> None:
    previous_path = tmp_path / "previous.xlsx"
    current_path = tmp_path / "current.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for row in range(1, 101):
        for column in range(1, 6):
            sheet.cell(row, column, row * column)
    workbook.save(previous_path)
    workbook.close()
    shutil.copyfile(previous_path, current_path)
    changed = load_workbook(current_path)
    changed.active["C50"] = 999_999
    changed.save(current_path)
    changed.close()

    incremental = ConsecutiveWorkbookReader()
    previous = incremental.read(previous_path)
    current = incremental.read(current_path)
    metrics = incremental.last_metrics

    assert current == read_workbook(current_path)
    assert metrics.rows_total == 100
    assert metrics.rows_reused == 99
    assert metrics.rows_parsed == 1
    assert metrics.cells_reused == 495
    assert metrics.cells_parsed == 5
    assert len(compare_snapshots(previous, current)) == 1


def test_dependency_change_conservatively_invalidates_every_row(tmp_path: Path) -> None:
    previous_path = tmp_path / "previous.xlsx"
    current_path = tmp_path / "current.xlsx"
    workbook = Workbook()
    for row in range(1, 11):
        workbook.active.cell(row, 1, row)
    workbook.save(previous_path)
    workbook.close()
    shutil.copyfile(previous_path, current_path)
    changed = load_workbook(current_path)
    changed.active["A1"].number_format = "0.0000"
    changed.save(current_path)
    changed.close()

    incremental = ConsecutiveWorkbookReader()
    incremental.read(previous_path)
    current = incremental.read(current_path)

    assert current == read_workbook(current_path)
    assert incremental.last_metrics.rows_reused == 0
    assert incremental.last_metrics.rows_parsed == 10


def test_row_fragments_inherit_all_realistic_worksheet_namespaces(tmp_path: Path) -> None:
    previous_path = tmp_path / "namespaces-before.xlsx"
    current_path = tmp_path / "namespaces-after.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = 1
    workbook.active["A2"] = 2
    workbook.save(previous_path)
    workbook.close()

    declarations = (
        b' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"'
        b' xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
        b' xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"'
        b' xmlns:x14ac="http://schemas.microsoft.com/office/spreadsheetml/2009/9/ac"'
        b' xmlns:xr="http://schemas.microsoft.com/office/spreadsheetml/2014/revision"'
        b' xmlns:custom="urn:trilha:test:unknown-valid-namespace"'
    )

    def add_realistic_namespaces(payload: bytes) -> bytes:
        payload = payload.replace(
            b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"',
            b'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
            + declarations,
            1,
        )
        payload = payload.replace(
            b'<row r="1">',
            b'<row r="1" r:id="relationship" mc:Ignorable="x14ac xr custom" '
            b'x14ac:dyDescent="0.25" xr:uid="{00000000-0000-0000-0000-000000000001}">',
        )
        payload = payload.replace(
            b'</row>',
            b'<x14:extension/><custom:metadata value="preserved"/></row>',
            1,
        )
        return payload

    _rewrite_member(
        previous_path, "xl/worksheets/sheet1.xml", add_realistic_namespaces
    )
    shutil.copyfile(previous_path, current_path)
    _rewrite_member(
        current_path,
        "xl/worksheets/sheet1.xml",
        lambda payload: payload.replace(b'<v>2</v>', b'<v>3</v>', 1),
    )

    incremental = ConsecutiveWorkbookReader()
    previous = incremental.read(previous_path)
    current = incremental.read(current_path)

    assert previous == read_workbook(previous_path)
    assert current == read_workbook(current_path)
    assert incremental.last_metrics.fallback_used is False
    assert incremental.last_metrics.rows_reused == 1
    assert incremental.last_metrics.rows_parsed == 1
    assert compare_snapshots(previous, current) == compare_snapshots(
        read_workbook(previous_path), read_workbook(current_path)
    )


def test_incremental_incompatibility_falls_back_and_reports_reason(
    monkeypatch: pytest.MonkeyPatch, cql028_versions: Path,
) -> None:
    def unsupported(*_args, **_kwargs):
        raise reader._FastReaderUnsupported("namespace experimental não suportado")

    monkeypatch.setattr(reader, "_iter_row_parts", unsupported)
    incremental = ConsecutiveWorkbookReader()
    path = cql028_versions / "0.85.xlsx"

    snapshot = incremental.read(path)

    assert snapshot == reader._read_openpyxl(path)
    assert incremental.last_metrics.fallback_used is True
    assert incremental.last_metrics.fallback_reason == (
        "namespace experimental não suportado"
    )
