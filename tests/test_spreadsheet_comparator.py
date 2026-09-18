import importlib.util
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

from app import spreadsheet_comparator as integrated


ALL_OPTIONS = {
    "value": True, "formula": True, "removed_content": True,
    "added_content": True, "font": True, "fill": True, "border": True,
    "alignment": True, "number_format": True, "protection": True,
    "added_rows": True, "removed_rows": True, "added_columns": True,
    "removed_columns": True, "row_height": True, "column_width": True,
    "row_hidden": True, "column_hidden": True, "merged_cells": True,
    "sheets": True, "sheet_state": True, "sheet_protection": True,
    "freeze_panes": True,
}


def _reference_module():
    path = Path(__file__).parents[1] / "reference" / "comparador_planilhas.py"
    spec = importlib.util.spec_from_file_location("reference_comparator", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader
    spec.loader.exec_module(module)
    return module


def _workbook(path: Path, sheets: dict[str, list[list[object]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    workbook.save(path)


@pytest.mark.parametrize("scenario", ["equal", "values", "sheets"])
def test_integrated_engine_has_reference_parity(tmp_path: Path, scenario: str) -> None:
    base = tmp_path / "base.xlsx"
    evaluated = tmp_path / "evaluated.xlsx"
    _workbook(base, {"Dados": [["ID", "Valor"], [1, 0], [2, False], [3, None]]})
    if scenario == "equal":
        _workbook(evaluated, {"Dados": [["ID", "Valor"], [1, 0], [2, False], [3, None]]})
    elif scenario == "values":
        _workbook(evaluated, {"Dados": [["ID", "Valor"], [1, "texto"], [2, 0], [4, None]]})
    else:
        _workbook(evaluated, {"Outra": [["ID"], [1]]})

    reference = _reference_module().compare_workbooks(base, evaluated, ALL_OPTIONS)
    actual = integrated.compare_workbooks(base, evaluated, ALL_OPTIONS)
    assert actual == reference


def test_invalid_or_missing_workbook_is_rejected(tmp_path: Path) -> None:
    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("not a workbook", encoding="utf-8")
    with pytest.raises(Exception):
        integrated.compare_workbooks(invalid, tmp_path / "missing.xlsx", ALL_OPTIONS)


def test_generated_result_contains_reference_content(tmp_path: Path) -> None:
    base = tmp_path / "base.xlsx"
    evaluated = tmp_path / "evaluated.xlsx"
    output = tmp_path / "result.xlsx"
    _workbook(base, {"Dados": [["ID", "Data"], [1, "2026-01-01"]]})
    _workbook(evaluated, {"Dados": [["ID", "Data"], [1, "2026-02-01"], [2, "novo"]]})
    differences = integrated.compare_workbooks(base, evaluated, ALL_OPTIONS)
    integrated.generate_report(str(base), str(evaluated), differences, str(output), ["Valor alterado"])
    workbook = load_workbook(output)
    try:
        assert workbook.sheetnames == ["Resumo", "Divergências"]
        assert workbook["Divergências"].max_row == len(differences) + 1
    finally:
        workbook.close()


def test_integrated_module_has_no_second_tk_or_mainloop() -> None:
    source = Path(integrated.__file__).read_text(encoding="utf-8")
    assert "tk.Tk(" not in source
    assert ".mainloop(" not in source
