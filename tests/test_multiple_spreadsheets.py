from pathlib import Path
import sqlite3

from openpyxl import Workbook, load_workbook

from app.audit_service import AuditService
from app.database import Database
from app.models import AuditExecutionStatus
from app.report_service import ReportService
from app.sources import LocalSource, SpreadsheetInfo, VersionInfo


def _write_workbook(path: Path, cells: dict[str, object]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Dados"
    for address, value in cells.items():
        worksheet[address] = value
    workbook.save(path)
    workbook.close()


def _history(
    directory: Path,
    prefix: str,
    definitions: list[tuple[str, dict[str, object]]],
) -> list[tuple[VersionInfo, Path]]:
    directory.mkdir()
    history = []
    for position, (number, cells) in enumerate(definitions, start=1):
        path = directory / f"{number}.xlsx"
        _write_workbook(path, cells)
        history.append(
            (
                VersionInfo(
                    # IDs deliberadamente iguais entre planilhas: o isolamento
                    # deve depender também da identidade técnica da planilha.
                    id=f"version-{position}",
                    number=number,
                    modified_at=f"2026-09-16T{position:02}:00:00Z",
                    author=f"Autor {prefix}",
                    comment=f"Versão {prefix} {number}",
                ),
                path,
            )
        )
    return history


def _identity(spreadsheet: SpreadsheetInfo) -> tuple[str, str, str]:
    return spreadsheet.site_id, spreadsheet.drive_id, spreadsheet.drive_item_id


def _counts_by_spreadsheet(
    connection: sqlite3.Connection, table: str
) -> dict[str, int]:
    rows = connection.execute(
        f"""SELECT p.drive_item_id, COUNT(*)
            FROM planilha p LEFT JOIN {table} item ON item.planilha_id = p.id
            GROUP BY p.id ORDER BY p.drive_item_id"""
    ).fetchall()
    return {row[0]: row[1] for row in rows}


def test_three_spreadsheets_share_one_database_without_state_mixing(
    tmp_path: Path,
) -> None:
    spreadsheets = (
        SpreadsheetInfo("site-a", "drive-a", "item-a", "Planilha A.xlsx", "/A.xlsx"),
        SpreadsheetInfo("site-b", "drive-b", "item-b", "Planilha B.xlsx", "/B.xlsx"),
        SpreadsheetInfo("site-c", "drive-c", "item-c", "Planilha C.xlsx", "/C.xlsx"),
    )
    sheet_a, sheet_b, sheet_c = spreadsheets
    history_a = _history(
        tmp_path / "a",
        "A",
        [
            ("1.0", {"A1": "A inicial"}),
            ("1.1", {"A1": "A consolidada"}),
            ("1.2", {"A1": "A final", "A2": "somente A"}),
        ],
    )
    history_b = _history(
        tmp_path / "b",
        "B",
        [
            ("2.0", {"B1": "B inicial"}),
            ("2.1", {"B1": "B parcial", "B2": "remover depois"}),
            ("2.2", {"B1": "B parcial", "B2": "remover depois", "C1": "somente B"}),
            ("2.3", {"B1": "B final", "C1": "somente B"}),
        ],
    )
    history_c = _history(
        tmp_path / "c",
        "C",
        [
            ("3.0", {}),
            ("3.1", {"C1": "C inicial"}),
            ("3.2", {"C1": "C final", "C2": "somente C"}),
        ],
    )

    database_path = tmp_path / "canonical.db"
    with Database(database_path) as database:
        database.initialize()

        # Estado controlado anterior: A totalmente consolidada, B somente até
        # 2.1 e C ainda inexistente no banco canônico.
        initial_source = LocalSource(
            spreadsheets,
            {
                _identity(sheet_a): history_a,
                _identity(sheet_b): history_b[:2],
                _identity(sheet_c): history_c,
            },
        )
        service = AuditService(database, initial_source)
        assert service.audit(sheet_a).final_version == "1.2"
        assert service.audit(sheet_b).final_version == "2.1"

        connection = database.connection
        assert _counts_by_spreadsheet(connection, "versao_processada") == {
            "item-a": 2,
            "item-b": 1,
        }
        assert _counts_by_spreadsheet(connection, "alteracao") == {
            "item-a": 3,
            "item-b": 2,
        }

        complete_source = LocalSource(
            spreadsheets,
            {
                _identity(sheet_a): history_a,
                _identity(sheet_b): history_b,
                _identity(sheet_c): history_c,
            },
        )
        acquired: list[tuple[str, str]] = []
        original_get_version = complete_source.get_version

        def record_acquisition(
            spreadsheet: SpreadsheetInfo, version: VersionInfo
        ) -> Path:
            acquired.append((spreadsheet.drive_item_id, version.number))
            return original_get_version(spreadsheet, version)

        complete_source.get_version = record_acquisition  # type: ignore[method-assign]
        service = AuditService(database, complete_source)

        no_updates = service.audit(sheet_a)
        incremental = service.audit(sheet_b)
        initial = service.audit(sheet_c)

        assert (
            no_updates.status,
            no_updates.processed_versions,
            no_updates.changes,
            no_updates.initial_checkpoint,
            no_updates.final_version,
        ) == (
            AuditExecutionStatus.COMPLETED_WITHOUT_UPDATES,
            0,
            0,
            "1.2",
            "1.2",
        )
        assert (
            incremental.status,
            incremental.processed_versions,
            incremental.changes,
            incremental.initial_checkpoint,
            incremental.final_version,
        ) == (AuditExecutionStatus.COMPLETED, 2, 3, "2.1", "2.3")
        assert (
            initial.status,
            initial.processed_versions,
            initial.changes,
            initial.initial_checkpoint,
            initial.final_version,
        ) == (AuditExecutionStatus.COMPLETED, 2, 3, None, "3.2")
        assert acquired == [
            ("item-b", "2.1"),
            ("item-b", "2.2"),
            ("item-b", "2.3"),
            ("item-c", "3.0"),
            ("item-c", "3.1"),
            ("item-c", "3.2"),
        ]

        checkpoints = {
            row[0]: (row[1], row[2])
            for row in connection.execute(
                """SELECT p.drive_item_id, c.versao_id, c.versao_numero
                   FROM checkpoint c JOIN planilha p ON p.id = c.planilha_id"""
            )
        }
        assert checkpoints == {
            "item-a": ("version-3", "1.2"),
            "item-b": ("version-4", "2.3"),
            "item-c": ("version-3", "3.2"),
        }
        assert _counts_by_spreadsheet(connection, "versao_processada") == {
            "item-a": 2,
            "item-b": 3,
            "item-c": 2,
        }
        assert _counts_by_spreadsheet(connection, "alteracao") == {
            "item-a": 3,
            "item-b": 5,
            "item-c": 3,
        }
        assert _counts_by_spreadsheet(connection, "execucao_auditoria") == {
            "item-a": 2,
            "item-b": 2,
            "item-c": 1,
        }

        # Nenhum relacionamento gravado pelo serviço atravessa identidades e
        # cada par, inclusive os IDs sobrepostos, ocorre exatamente uma vez.
        assert connection.execute(
            """SELECT COUNT(*) FROM alteracao a
               JOIN versao_processada v ON v.id = a.versao_processada_id
               WHERE a.planilha_id <> v.planilha_id"""
        ).fetchone()[0] == 0
        assert connection.execute(
            """SELECT COUNT(*) FROM versao_processada v
               JOIN execucao_auditoria e ON e.id = v.execucao_id
               WHERE v.planilha_id <> e.planilha_id"""
        ).fetchone()[0] == 0
        assert connection.execute(
            """SELECT COUNT(*) FROM (
                   SELECT planilha_id, versao_anterior_id, versao_atual_id, COUNT(*) total
                   FROM versao_processada
                   GROUP BY planilha_id, versao_anterior_id, versao_atual_id
                   HAVING total > 1
               )"""
        ).fetchone()[0] == 0

        expected_reports = {
            "item-a": ({"1.0", "1.1", "1.2"}, 3),
            "item-b": ({"2.0", "2.1", "2.2", "2.3"}, 5),
            "item-c": ({"3.0", "3.1", "3.2"}, 3),
        }
        report_service = ReportService(connection, tmp_path / "reports")
        for spreadsheet in spreadsheets:
            spreadsheet_id = connection.execute(
                """SELECT id FROM planilha
                   WHERE site_id = ? AND drive_id = ? AND drive_item_id = ?""",
                _identity(spreadsheet),
            ).fetchone()[0]
            report_path = report_service.generate(spreadsheet_id)
            workbook = load_workbook(report_path, data_only=False)
            try:
                summary = dict(
                    workbook["RESUMO"].iter_rows(min_row=2, values_only=True)
                )
                versions = {
                    value
                    for row in workbook["VERSOES"].iter_rows(
                        min_row=2, values_only=True
                    )
                    for value in row[:2]
                }
                expected_versions, expected_changes = expected_reports[
                    spreadsheet.drive_item_id
                ]
                assert summary["DriveItem ID"] == spreadsheet.drive_item_id
                assert summary["Total de alterações"] == expected_changes
                assert versions == expected_versions
                assert workbook["TRILHA"].max_row == expected_changes + 1
            finally:
                workbook.close()

        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []
